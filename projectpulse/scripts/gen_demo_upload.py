"""Write a three-project program as workbooks somebody can upload live.

    python -m scripts.gen_demo_upload                 # into data/upload_demo/
    python -m scripts.gen_demo_upload --out somewhere

Six files - a schedule and a worklog per project - in the app's own template
shape, so they go in through `Settings > Sources` exactly as a PM's own sheets
would. Nothing here seeds a domain row: the real reader, identity resolver,
differ and schedule engine run over these the same as over anything else, which
is the only way a demo proves the product rather than a screenshot.

**Why not a Jira export.** A Jira export carries no baseline, no dependency
edges and no effort (see `app/ingest/sources/jira/export_sheet.py`), so a
project imported from one can say what is overdue and what is piling up but
never *what it will cost* - there is no chain for slip to propagate along. This
set exists for the other half of the story, so the schedule has something to be
chaotic about.

**What makes it messy on purpose.** A demo where every row is clean proves
nothing about a hand-maintained plan, so each project carries a different
failure:

``PAY`` Payments Core
    A long chain with an early slip that propagates the whole way down, and two
    tasks dated *earlier than their own predecessors allow* - the inconsistency
    a sheet cannot show you itself.
``POR`` Customer Portal
    Waits on Payments through a cross-project hand-off it cannot see, and its QA
    queue has stalled: most items blocked, aging, several on the customer.
``MIG`` Data Migration
    Its schedule is genuinely clean - every date met, no inconsistency, nothing
    propagating - and it is still not fine: 61 hours logged against 23 planned.
    The contrast is the point. A portfolio where everything is red teaches a
    reader to ignore the colour, and a project that is green on schedule and
    amber on effort is the case a single health score would flatten away.

Owners are shared across all three deliberately. Contention itself needs
`Resource` rows, which arrive through `scripts.seed_extras` rather than an
upload, but the names lining up is what makes that story legible when it does.

Verified by uploading the whole set through `POST /api/sources/upload` into a
scratch database: six files, **zero rejected rows**, and the three projects come
out `critical` / `critical` / `watch` with different band patterns - which is the
gradient that makes a heat-map worth looking at.
"""

from __future__ import annotations

# Must run before any `app.*` import - see scripts/_bootstrap.py.
from scripts._bootstrap import bootstrap

bootstrap()

import argparse  # noqa: E402
from datetime import date, timedelta  # noqa: E402
from pathlib import Path  # noqa: E402

from openpyxl import Workbook  # noqa: E402

from app.ingest.sources.excel.reader import (  # noqa: E402
    SCHEDULE_CONTRACT,
    WORKLOG_CONTRACT,
)

#: The story's "today". Everything is dated relative to it so the set stays
#: meaningful whenever it is generated - a fixed calendar would drift into the
#: past and every task would read as overdue, which is the trap `tasks_overdue`
#: is calibrated around.
TODAY = date.today()


def d(offset: int) -> date:
    return TODAY + timedelta(days=offset)


#: (task id, activity, phase, milestone, status, owner, start, baseline, planned,
#:  progress, predecessor)
#:
#: Offsets rather than dates, so the shape of the plan is readable as a shape.
PROJECTS: dict[str, dict] = {
    "PAY": {
        "name": "Payments Core",
        "schedule": [
            ("PAY-1", "Requirements sign-off", "Requirement", "Foundations", "Done", "Nguyen Van A", -40, -30, -30, 100, ""),
            ("PAY-2", "Ledger schema design", "Design", "Foundations", "Done", "Le Van C", -30, -20, -16, 100, "PAY-1"),
            # The slip that starts everything: committed -12, actually landing +4.
            ("PAY-3", "Settlement engine", "Development", "Core build", "In Progress", "Tran Quoc B", -16, -12, 4, 55, "PAY-2"),
            ("PAY-4", "Reconciliation service", "Development", "Core build", "In Progress", "Nguyen Van A", -8, -4, 10, 30, "PAY-3"),
            # Dated before its predecessor can finish - the inconsistency a
            # spreadsheet cannot show you.
            ("PAY-5", "Fraud rules integration", "Development", "Core build", "Not started", "Le Van C", 0, 6, 6, 0, "PAY-4"),
            ("PAY-6", "Payment gateway cutover", "Deployment", "Go-live", "Not started", "Tran Quoc B", 8, 16, 16, 0, "PAY-5"),
            # Same again, harder: a go-live milestone dated before the work under it.
            ("PAY-7", "Production smoke tests", "Testing", "Go-live", "Not started", "My Nguyen", 14, 18, 18, 0, "PAY-6"),
            ("PAY-8", "Regulator evidence pack", "Requirement", "Go-live", "Not started", "Pham Hong D", 10, 20, 20, 0, "PAY-4"),
        ],
        "worklog": [
            ("PAY-Q1", "Rounding on multi-currency settlement", "Open", "Yes", "Tran Quoc B", 12, 14, -18),
            ("PAY-Q2", "Ledger reconciliation mismatch", "Open", "Yes", "Nguyen Van A", 8, 11, -15),
            ("PAY-Q3", "Retry storm under gateway timeout", "Open", "Yes", "Tran Quoc B", 6, 9, -11),
            ("PAY-Q4", "Fee schedule confirmation", "Open", "Customer", "Pham Hong D", 3, 1, -9),
            ("PAY-Q5", "Audit log retention window", "Open", "No", "Le Van C", 4, 4, -6),
            ("PAY-Q6", "Settlement batch performance", "Closed", "No", "Tran Quoc B", 10, 10, -20),
        ],
    },
    "POR": {
        "name": "Customer Portal",
        "schedule": [
            ("POR-1", "Portal UX sign-off", "Design", "Foundations", "Done", "My Nguyen", -35, -26, -26, 100, ""),
            ("POR-2", "Account service API", "Development", "Foundations", "Done", "Le Van C", -26, -16, -14, 100, "POR-1"),
            ("POR-3", "Statements screen", "Development", "Portal build", "In Progress", "My Nguyen", -14, -6, 2, 60, "POR-2"),
            # Waits on the payments hand-off, which this project cannot see.
            ("POR-4", "Payment history integration", "Development", "Portal build", "Not started", "Le Van C", -2, 4, 12, 0, "POR-3"),
            ("POR-5", "Accessibility audit", "Testing", "Portal build", "Not started", "Pham Hong D", 6, 12, 12, 0, "POR-3"),
            ("POR-6", "Customer UAT", "Testing", "Release", "Not started", "My Nguyen", 12, 18, 22, 0, "POR-4"),
            ("POR-7", "Portal release", "Deployment", "Release", "Not started", "Tran Quoc B", 20, 24, 24, 0, "POR-6"),
        ],
        "worklog": [
            ("POR-Q1", "Statement PDF layout on mobile", "Open", "Yes", "My Nguyen", 5, 8, -12),
            ("POR-Q2", "Session timeout on slow networks", "Open", "Yes", "Le Van C", 4, 7, -10),
            ("POR-Q3", "Which balances are shown pre-settlement", "Open", "Customer", "Pham Hong D", 2, 0, -14),
            ("POR-Q4", "Contrast ratio on the statements table", "Open", "Customer", "My Nguyen", 3, 1, -13),
            ("POR-Q5", "Locale formatting for amounts", "Open", "No", "Le Van C", 3, 3, -5),
            ("POR-Q6", "Login redirect loop", "Closed", "No", "Tran Quoc B", 6, 6, -19),
        ],
    },
    "MIG": {
        "name": "Data Migration",
        "schedule": [
            ("MIG-1", "Source system profiling", "Requirement", "Discovery", "Done", "Pham Hong D", -38, -28, -28, 100, ""),
            ("MIG-2", "Mapping specification", "Design", "Discovery", "Done", "Hoach Bach", -28, -18, -18, 100, "MIG-1"),
            ("MIG-3", "Extract pipeline", "Development", "Build", "Done", "Pham Hong D", -18, -8, -8, 100, "MIG-2"),
            ("MIG-4", "Transform rules", "Development", "Build", "In Progress", "Hoach Bach", -8, 2, 2, 40, "MIG-3"),
            ("MIG-5", "Dry-run load", "Testing", "Rehearsal", "Not started", "Pham Hong D", 4, 12, 12, 0, "MIG-4"),
            ("MIG-6", "Cutover rehearsal", "Testing", "Rehearsal", "Not started", "Hoach Bach", 14, 22, 22, 0, "MIG-5"),
        ],
        "worklog": [
            ("MIG-Q1", "Null handling in legacy address rows", "Closed", "No", "Pham Hong D", 6, 14, -22),
            ("MIG-Q2", "Duplicate customer keys", "Closed", "No", "Hoach Bach", 8, 19, -20),
            ("MIG-Q3", "Timezone drift on created dates", "Open", "No", "Pham Hong D", 4, 12, -16),
            ("MIG-Q4", "Charset on legacy notes", "Open", "No", "Hoach Bach", 5, 16, -14),
        ],
    },
}


def _sheet(workbook: Workbook, title: str, headers) -> None:
    sheet = workbook.active
    sheet.title = title
    sheet.append(list(headers))
    for index, header in enumerate(headers, start=1):
        column = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[column].width = max(12, len(header) + 2)


def write_schedule(key: str, out: Path) -> Path:
    workbook = Workbook()
    _sheet(workbook, "Activities", SCHEDULE_CONTRACT.template_headers)
    sheet = workbook.active
    for row in PROJECTS[key]["schedule"]:
        task, activity, phase, milestone, status, owner, start, base, plan, prog, pred = row
        sheet.append(
            [task, activity, phase, milestone, status, owner, d(start), d(base),
             d(plan), prog, pred, d(-2)]
        )
    path = out / f"{key.lower()}_schedule.xlsx"
    workbook.save(path)
    return path


def write_worklog(key: str, out: Path) -> Path:
    workbook = Workbook()
    _sheet(workbook, "Worklog", WORKLOG_CONTRACT.template_headers)
    sheet = workbook.active
    for task, summary, status, blocked, owner, estimate, hours, logged in PROJECTS[key]["worklog"]:
        sheet.append([task, summary, status, blocked, owner, estimate, hours, d(logged)])
    path = out / f"{key.lower()}_worklog.xlsx"
    workbook.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--out", type=Path, default=Path("data") / "upload_demo",
        help="where to write the workbooks",
    )
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    written = []
    for key, spec in PROJECTS.items():
        written.append((spec["name"], write_schedule(key, args.out), write_worklog(key, args.out)))

    print(f"wrote {len(written) * 2} workbook(s) to {args.out}\n")
    for name, schedule, worklog in written:
        print(f"  {name}")
        print(f"    schedule  {schedule.name}")
        print(f"    worklog   {worklog.name}")
    print(
        "\n  Upload each pair at Settings > Sources - kind 'schedule' then 'worklog',\n"
        "  same project name for both, and the same program for all three so the\n"
        "  cross-project rollup has something to roll up.\n"
    )


if __name__ == "__main__":
    main()
