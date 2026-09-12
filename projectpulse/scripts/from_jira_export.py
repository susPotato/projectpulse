"""Turn a Jira issue export into the schedule workbook this app reads.

    python -m scripts.from_jira_export "Jira Cowork Local 1.xlsx"
    python -m scripts.from_jira_export export.xlsx --out data/coworklocal_schedule.xlsx

Jira's "general_report" export is not a sheet this app can read, and the gap is
not only the column names. Three things differ:

**Its table does not start at row 1.** There is a filter name, a "Displaying N
issues at ..." line, and then the header. `reader.find_sheet` already searches
for a header row rather than assuming one, so that half would have worked.

**One issue is not one row.** An export with a rich-text Description writes each
issue across a block of rows, the fields on the first and the wrapped text
below. Rows are therefore collected by "has a Key", never by position.

**Four hundred columns, and the ones this app needs are mostly not there.** The
export carries every custom field the Jira instance defines - most of them
empty, several containing the page's own JavaScript rather than a value. What
it does *not* carry is a baseline, and that is the important absence: see below.

What this writes is the blank schedule template's own column list
(`SCHEDULE_CONTRACT.template_headers`), so the output is the same document a
person would get from `GET /api/template/schedule.xlsx` and fill in by hand.
Upload it at Settings > Sources, or POST it to `/api/sources/upload` with
`sheet_kind=schedule`.

**It fills the columns the export actually has, and leaves the rest empty
rather than inventing them.** That is the whole design of this script, and the
reason it prints a coverage report instead of just succeeding: a Jira export of
this shape cannot light up most of this product, and finding that out from a
blank dashboard is worse than being told up front.

- **No baseline.** Jira has no "the date this was originally committed to"
  field, so `Baseline Finish` is empty and every recorded-slip figure is 0.
  Slip a human typed is not knowable from a Jira export - only from two exports
  taken at different times, which is what the differ would compare.
- **No dependency edges** unless `Linked Issues` or `Sub-Tasks` are populated.
  This is the one that matters most: the projected finish, propagated slip, the
  driving path, `days_late` and milestones-at-risk are all forward-pass results
  over a DAG. With no edges, projected == planned for every task and the
  intelligence layer has nothing to say.
- **No effort** unless `Original Estimate` / `Time Spent` are populated, and
  those belong in a *worklog* sheet, not this one.
"""

from __future__ import annotations

# Must run before any `app.*` import - see scripts/_bootstrap.py.
from scripts._bootstrap import bootstrap

bootstrap()

import argparse  # noqa: E402
import re  # noqa: E402
from datetime import date, datetime  # noqa: E402
from pathlib import Path  # noqa: E402

from openpyxl import Workbook, load_workbook  # noqa: E402

from app.ingest.sources.excel.reader import SCHEDULE_CONTRACT  # noqa: E402

#: How far down to look for the header row. Generous because the preamble is a
#: report title and a "Displaying N issues" line today, and nothing promises a
#: different filter will not add a third.
HEADER_SEARCH_ROWS = 25

#: Which Jira column feeds each column of the schedule template, most-preferred
#: first. Matched case-insensitively on the header text.
#:
#: Several entries list more than one name because "the" Jira column depends on
#: how the instance was configured: `Due Date` is the system field and `End
#: date` is the custom one an Advanced-Roadmaps instance uses instead, and a
#: given export usually has exactly one of them filled.
SOURCES: dict[str, tuple[str, ...]] = {
    "Task ID": ("Key",),
    "Activity": ("Summary",),
    "Phase": ("Issue Type",),
    "Milestone": ("Milestone", "Fix Version/s"),
    "Status": ("Status",),
    "Owner": ("Assignee",),
    "Start": ("Planned Start", "Start date", "Start Date"),
    # Deliberately fed by nothing - Jira has no baseline field. Listed anyway so
    # the coverage report names it rather than leaving a reader to notice.
    "Baseline Finish": (),
    "Planned Finish": ("Due Date", "End date"),
    "Progress": ("Task progress", "Progress"),
    "Predecessor": ("Linked Issues", "Sub-Tasks"),
}

#: Jira renders some custom fields by shipping the page's own script, so the
#: cell holds a function body rather than a value. Any cell that looks like this
#: is dropped rather than written through - a task whose Owner is a jQuery call
#: is worse than a task with no Owner.
SCRIPT_SMELL = re.compile(r"\$\(|function\s*\(|setTimeout|document\.ready", re.I)

#: Same idea for Jira's own rendering failures, which arrive as prose.
ERROR_SMELL = re.compile(r"^Error rendering ", re.I)


def _clean(value):
    """One cell, as a value or not at all."""
    if value is None:
        return None
    if isinstance(value, (int, float, date, datetime)):
        return value
    text = str(value).strip()
    if not text or text.casefold() in {"none", "unresolved", "n/a"}:
        return None
    if SCRIPT_SMELL.search(text) or ERROR_SMELL.match(text):
        return None
    return text


def _as_date(value):
    """A date, or None. Jira writes datetimes; a schedule column wants a day.

    The time is dropped rather than kept because this app compares dates, and
    carrying `19:22` into a planned-finish column would suggest the plan is
    precise to the minute when the PM chose a day.
    """
    value = _clean(value)
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    try:
        return datetime.fromisoformat(str(value)[:10]).date()
    except ValueError:
        return None


def _find_header(grid: list[tuple]) -> int | None:
    """The row holding the export's column names.

    Identified by containing `Key` *and* `Summary`, not by position and not by
    `Key` alone - a preamble line can easily contain one word.
    """
    for index, row in enumerate(grid):
        labels = {str(c).strip().casefold() for c in row if c is not None}
        if "key" in labels and "summary" in labels:
            return index
    return None


def read_export(path: Path, sheet: str | None = None) -> tuple[list[str], list[tuple]]:
    """`(headers, issue_rows)` from a Jira export, whichever tab holds it."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        names = ([sheet] if sheet else []) + [n for n in workbook.sheetnames if n != sheet]
        for name in names:
            worksheet = workbook[name]
            grid = []
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= HEADER_SEARCH_ROWS:
                    break
                grid.append(row)
            header_at = _find_header(grid)
            if header_at is None:
                continue

            headers = [("" if c is None else str(c).strip()) for c in grid[header_at]]
            key_at = [h.casefold() for h in headers].index("key")

            rows = []
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index <= header_at:
                    continue
                # One issue is one row *with a Key*. Everything between is the
                # previous issue's wrapped rich text.
                if key_at < len(row) and _clean(row[key_at]):
                    rows.append(row)
            return headers, rows
    finally:
        workbook.close()
    raise SystemExit(f"no Jira export table found in {path} - no header row with Key + Summary")


def convert(headers: list[str], rows: list[tuple]) -> tuple[list[dict], dict[str, str]]:
    """`(records, chosen)` - the template rows, and which Jira column fed each."""
    lookup = {h.casefold(): i for i, h in enumerate(headers) if h}

    chosen: dict[str, str] = {}
    for column, candidates in SOURCES.items():
        for candidate in candidates:
            if candidate.casefold() in lookup:
                chosen[column] = candidate
                break

    date_columns = {"Start", "Planned Finish", "Baseline Finish"}
    records = []
    for row in rows:
        record = {}
        for column in SCHEDULE_CONTRACT.template_headers:
            source = chosen.get(column)
            if source is None:
                record[column] = None
                continue
            raw = row[lookup[source.casefold()]] if lookup[source.casefold()] < len(row) else None
            record[column] = _as_date(raw) if column in date_columns else _clean(raw)
        records.append(record)
    return records, chosen


def write_schedule(records: list[dict], out: Path) -> None:
    """The blank template's own shape, filled in.

    The tab is named `Activities` because that is what the contract prefers and
    `find_sheet` tries first - it would resolve a differently-named tab too, but
    only by searching, and there is no reason to make it search.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Activities"
    sheet.append(list(SCHEDULE_CONTRACT.template_headers))
    for record in records:
        sheet.append([record[h] for h in SCHEDULE_CONTRACT.template_headers])
    for index, header in enumerate(SCHEDULE_CONTRACT.template_headers, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            12, len(header) + 2
        )
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)


def report(records: list[dict], chosen: dict[str, str], headers: list[str], rows: list[tuple]) -> None:
    """What came across, what did not, and what that costs.

    Printed every run rather than behind a flag: a conversion that "succeeded"
    into a schedule with no baseline and no edges is the case worth being loud
    about, and it is also the likely one.
    """
    total = len(records)
    print(f"\n  {total} issue(s) -> {len(SCHEDULE_CONTRACT.template_headers)} schedule columns\n")
    for column in SCHEDULE_CONTRACT.template_headers:
        filled = sum(1 for r in records if r[column] not in (None, ""))
        source = chosen.get(column)
        origin = f"from {source!r}" if source else "no Jira column feeds this"
        flag = "  " if filled else "! "
        print(f"  {flag}{column:16} {filled:>3}/{total}  {origin}")

    missing = [c for c in SCHEDULE_CONTRACT.template_headers
               if not any(r[c] not in (None, "") for r in records)]

    if "Baseline Finish" in missing:
        print(
            "\n  ! No baseline. Jira has no 'originally committed to' field, so every\n"
            "    recorded-slip figure will read 0. Slip a person typed is only\n"
            "    recoverable from two exports taken at different times."
        )
    if "Predecessor" in missing:
        print(
            "  ! No dependency edges (Linked Issues / Sub-Tasks are empty). The\n"
            "    projected finish, propagated slip, the driving path, days-late and\n"
            "    milestones-at-risk are all forward-pass results over a DAG. Without\n"
            "    edges every task's projected date equals its planned one, and the\n"
            "    findings that make this product worth opening will not fire."
        )

    # A custom field that defaults to the creation stamp is not a plan, and it
    # is indistinguishable from one once it is in a Start column.
    start_source = chosen.get("Start")
    if start_source and "created" in {h.casefold() for h in headers}:
        lookup = {h.casefold(): i for i, h in enumerate(headers) if h}
        s, c = lookup[start_source.casefold()], lookup["created"]
        # Compared only over rows that *have* a start. Requiring every row to
        # match meant one issue with the field left blank hid the collision on
        # the other sixteen, which is the wrong way round: a blank is not
        # evidence that the populated ones are real.
        pairs = [(r[s], r[c]) for r in rows
                 if s < len(r) and c < len(r) and r[s] is not None and r[c] is not None]
        same = sum(1 for a, b in pairs if a == b)
        if pairs and same == len(pairs):
            print(
                f"\n  ! {start_source!r} equals 'Created' on all {len(pairs)} rows that have it.\n"
                "    That is a field default, not a planned start - it was written when\n"
                "    the issue was raised. It has been carried across because it is what\n"
                "    the export says, but do not read the Start column as a plan."
            )

    distinct_owners = {r["Owner"] for r in records if r["Owner"]}
    if len(distinct_owners) <= 1:
        print(
            f"\n  ! {len(distinct_owners)} distinct owner(s). Resource contention is computed\n"
            "    across projects sharing a person, so a single-owner project contributes\n"
            "    nothing to a program rollup until other projects name the same people."
        )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("export", type=Path, help="the Jira .xlsx export")
    parser.add_argument("--out", type=Path, default=None, help="where to write the schedule workbook")
    parser.add_argument("--sheet", default=None, help="which tab of the export holds the table")
    args = parser.parse_args()

    if not args.export.exists():
        raise SystemExit(f"no such file: {args.export}")

    headers, rows = read_export(args.export, args.sheet)
    records, chosen = convert(headers, rows)
    if not records:
        raise SystemExit("the export has a header row but no issues under it")

    out = args.out or args.export.with_name(args.export.stem + "_schedule.xlsx")
    write_schedule(records, out)
    print(f"wrote {out}")
    report(records, chosen, headers, rows)
    print(
        "\n  Load it: Settings > Sources > upload, kind 'schedule'. Or:\n"
        f"    curl -F file=@'{out}' -F sheet_kind=schedule -F project_name='<name>' \\\n"
        "         http://127.0.0.1:8000/api/sources/upload\n"
    )


if __name__ == "__main__":
    main()
