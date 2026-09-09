"""The status report a project manager circulates, as a Word document.

The same findings as the Insight screen, in the format the conversation around
them actually happens in - a document attached to an email, read by people who
will never open the app.

**It decides nothing about content.** The report is built by
`app/exports/document.py` into blocks, and this module only turns a block into a
Word paragraph. That split is what lets the same report come out as `.xlsx` and
Markdown without three exporters drifting apart, and it means a number cannot be
formatted here: every string arriving in a block was already formatted by
`assembler.format_fact` or is already-substituted prose from `Finding.headline`.
Invariant 1 holds on the artefact that outlives the session - a figure rendered
its own way here would be the copy quoted in a steering meeting six weeks later,
disagreeing with a screen nobody has open.

`python-docx` is an optional extra. Nothing else imports this module, so the
demo, the tests and the API all run without it - `report_bytes` raises a plain
message naming the extra if it is missing. The other two formats need no extra
at all, which is why the builder screen does not offer Word as its only choice.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Sequence

from app.api.schemas.explain import ExplainBundle
from app.api.schemas.insight import InsightBundle
from app.exports.document import (
    MAX_EVIDENCE,
    MAX_PROJECTION_ROWS,
    Block,
    ReportDoc,
    build_document,
)

__all__ = [
    "MAX_EVIDENCE",
    "MAX_PROJECTION_ROWS",
    "ReportUnavailable",
    "build_report",
    "render_docx",
    "report_bytes",
]

#: Word's own grid style. Named rather than hand-drawn so the table inherits the
#: reader's theme instead of carrying ours into their house style.
_TABLE_STYLE = "Light Grid Accent 1"


class ReportUnavailable(RuntimeError):
    """`python-docx` is not installed."""


def _docx():
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise ReportUnavailable(
            'python-docx is not installed; pip install -e ".[report]"'
        ) from exc
    return docx


def _add_block(document, block: Block) -> None:
    """One block, in Word's vocabulary."""
    if block.kind == "heading":
        document.add_heading(block.text, level=block.level)
        return

    if block.kind == "paragraph":
        if not block.label:
            document.add_paragraph(block.text)
            return
        # Two runs rather than one string, so the label is bold and the sentence
        # is not. Splitting a finished sentence back into label and body would
        # mean guessing where the colon was.
        paragraph = document.add_paragraph()
        paragraph.add_run(block.label).bold = True
        if block.text:
            paragraph.add_run(block.text)
        return

    if block.kind == "note":
        paragraph = document.add_paragraph()
        paragraph.add_run(block.text).italic = True
        return

    if block.kind == "bullets":
        for item in block.items:
            document.add_paragraph(item, style="List Bullet")
        return

    if block.kind == "table":
        if block.text:
            document.add_paragraph(block.text)
        table = document.add_table(rows=1, cols=len(block.columns))
        table.style = _TABLE_STYLE
        for cell, label in zip(table.rows[0].cells, block.columns):
            cell.text = label
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        for row in block.rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
        return


def render_docx(doc: ReportDoc):
    """A built `ReportDoc` as a `docx.Document`, ready to save."""
    docx = _docx()
    document = docx.Document()

    document.add_heading(doc.title, 0)
    for block in doc.preamble:
        _add_block(document, block)

    for section in doc.sections:
        document.add_heading(section.title, level=1)
        for block in section.blocks:
            _add_block(document, block)

    return document


def build_report(
    bundle: InsightBundle,
    *,
    explain: ExplainBundle | None = None,
    scenarios=None,
    risks=None,
    sections: Sequence[str] | None = None,
    project_name: str = "",
    generated_at: datetime | None = None,
):
    """The whole report as a `docx.Document`, ready to save."""
    return render_docx(
        build_document(
            bundle,
            explain=explain,
            scenarios=scenarios,
            risks=risks,
            sections=sections,
            project_name=project_name,
            generated_at=generated_at,
        )
    )


def report_bytes(
    bundle: InsightBundle,
    *,
    explain: ExplainBundle | None = None,
    scenarios=None,
    risks=None,
    sections: Sequence[str] | None = None,
    project_name: str = "",
    generated_at: datetime | None = None,
) -> bytes:
    """The report as bytes, for an HTTP response or a file write."""
    buffer = BytesIO()
    build_report(
        bundle,
        explain=explain,
        scenarios=scenarios,
        risks=risks,
        sections=sections,
        project_name=project_name,
        generated_at=generated_at,
    ).save(buffer)
    return buffer.getvalue()
