"""Best-effort URL reading, so a pasted link behaves like an attachment.

Same idea as `pimsathon-main`'s `core/link_fetch.py`, reimplemented on the
stdlib plus what this project already depends on (`openpyxl`) or optionally
installs (`python-docx`, the `report` extra) - no new hard dependency for a
feature that is itself optional. PDF is deliberately not supported: parsing
it needs a new package this project has no other use for, and the honest
failure ("could not read a PDF at this link yet") costs nothing a silent gap
would not have cost anyway.

Stdlib `urllib` rather than `requests`, on purpose: `requests` is only
installed with the `onedrive` extra, and URL reading should not require
signing in to OneDrive to work.
"""

from __future__ import annotations

import re
import tempfile
import urllib.request
from html.parser import HTMLParser
from pathlib import Path
from urllib.error import URLError

#: Plenty for a model to work with; more than this just spends tokens on
#: content the reply will not use anyway.
MAX_CHARS = 20_000
MAX_DOWNLOAD_BYTES = 15 * 1024 * 1024
TIMEOUT_SECONDS = 15

_URL_RE = re.compile(r"https?://[^\s<>\"']+")


def find_first_url(text: str) -> str | None:
    match = _URL_RE.search(text)
    return match.group(0) if match else None


class _TextCollector(HTMLParser):
    """The crudest HTML-to-text there is: drop tags, keep what a reader sees."""

    _SKIP = {"script", "style", "noscript"}

    def __init__(self) -> None:
        super().__init__()
        self._skip_depth = 0
        self.chunks: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip_depth += 1

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data):
        if not self._skip_depth and data.strip():
            self.chunks.append(data.strip())


def _extract_html(raw: bytes) -> str:
    parser = _TextCollector()
    parser.feed(raw.decode("utf-8", errors="replace"))
    return "\n".join(parser.chunks)


def _write_temp(raw: bytes, suffix: str) -> Path:
    """A closed temp file holding `raw` - Windows refuses to let a second
    handle (openpyxl's, python-docx's) open a `NamedTemporaryFile` that is
    still held open by this one, so the file has to be closed before it is
    read back and the caller is responsible for deleting it."""
    handle = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    try:
        handle.write(raw)
    finally:
        handle.close()
    return Path(handle.name)


def _extract_xlsx(raw: bytes) -> str:
    import openpyxl

    path = _write_temp(raw, ".xlsx")
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"# Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    if any(cell is not None for cell in row):
                        lines.append(", ".join("" if c is None else str(c) for c in row))
            return "\n".join(lines)
        finally:
            # `read_only` streams from the zip lazily and keeps it open until
            # closed explicitly - on Windows the `unlink` below fails with
            # "used by another process" otherwise.
            workbook.close()
    finally:
        path.unlink(missing_ok=True)


def _extract_docx(raw: bytes) -> str | None:
    try:
        import docx
    except ImportError:
        return None

    path = _write_temp(raw, ".docx")
    try:
        document = docx.Document(path)
        return "\n".join(p.text for p in document.paragraphs if p.text.strip())
    finally:
        path.unlink(missing_ok=True)


def fetch_and_extract(url: str) -> str:
    """The text at `url`, or a one-line note explaining why there isn't any.

    Never raises - a failed fetch is content for the model to see and
    mention, the same way a missing sheet is a note rather than a crash
    elsewhere in this app.
    """
    try:
        request = urllib.request.Request(url, headers={"User-Agent": "ProjectPulseAI-Agent/1.0"})
        with urllib.request.urlopen(request, timeout=TIMEOUT_SECONDS) as response:
            content_type = (response.headers.get("Content-Type") or "").lower()
            raw = response.read(MAX_DOWNLOAD_BYTES + 1)
    except URLError as exc:
        return f"[could not fetch {url}: {exc}]"
    except Exception as exc:  # noqa: BLE001 - any failure becomes a visible note
        return f"[could not fetch {url}: {type(exc).__name__}: {exc}]"

    if len(raw) > MAX_DOWNLOAD_BYTES:
        return f"[{url} is too large to read - skipped]"

    lower_url = url.lower()
    try:
        if "text/html" in content_type or lower_url.endswith((".html", ".htm")):
            text = _extract_html(raw)
        elif "spreadsheet" in content_type or lower_url.endswith((".xlsx", ".xlsm")):
            text = _extract_xlsx(raw)
        elif lower_url.endswith(".docx") or "wordprocessingml" in content_type:
            text = _extract_docx(raw)
            if text is None:
                return (
                    f"[{url} is a .docx - install the 'report' extra "
                    '(pip install -e ".[report]") to read it]'
                )
        elif lower_url.endswith(".pdf") or "application/pdf" in content_type:
            return f"[{url} is a PDF - reading PDFs is not supported yet]"
        else:
            # Best effort: treat as plain text (markdown, .txt, .csv, .json...).
            text = raw.decode("utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001 - extraction failing is still a note
        return f"[could not read {url}: {type(exc).__name__}: {exc}]"

    text = text.strip()
    if not text:
        return f"[{url} had no readable text]"
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n[truncated]"
    return text
