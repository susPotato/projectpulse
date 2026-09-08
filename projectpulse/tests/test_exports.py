"""Guards for the two files the product hands back.

The template tests are round-trips rather than assertions about bytes. A
template that opens in Excel and that the ingester cannot read is worse than no
template at all - it fails at the point where a PM has already done the work of
filling it in - so the test writes one and reads it with the real reader.

The report tests are about invariant 1. A report is the artefact that outlives
the session, so a number formatted its own way here is the copy that gets quoted
in a steering meeting six weeks later.
"""

from __future__ import annotations

import importlib
import tempfile
from datetime import date, datetime, timezone
from pathlib import Path

import pytest

from app.api.schemas.explain import ExplainBundle, ForwardStep
from app.api.schemas.insight import (
    DataQuality,
    EvidenceRef,
    Finding,
    InsightBundle,
    RuleTrace,
)
from app.exports.template import KINDS, template_bytes, template_workbook
from app.ingest.sources.excel.reader import (
    SCHEDULE_CONTRACT,
    WORKLOG_CONTRACT,
    read_sheet,
)

# --------------------------------------------------------------------------
# The blank template
# --------------------------------------------------------------------------


@pytest.mark.parametrize("kind", sorted(KINDS))
def test_a_generated_template_is_readable_by_the_real_ingester(kind):
    """The test the whole module exists for.

    Not "does openpyxl produce a file" but "does our own reader find the header
    row, recognise the key column, and report no unknown headers". Anything less
    and a PM fills in a sheet we then refuse.
    """
    contract, sheet_name, _ = KINDS[kind]

    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / f"{kind}.xlsx"
        path.write_bytes(template_bytes(kind))
        read = read_sheet(path, sheet_name, contract)

    assert read.header_row is not None, "the reader could not find the headers"
    assert read.has_key_column, "no Task ID column, so no row can be tracked"
    assert read.unknown_headers == [], f"headers we write but cannot read: {read.unknown_headers}"
    assert read.rejects == []
    # Blank by design - an example row would come back as a real task.
    assert read.rows == []


@pytest.mark.parametrize("kind", sorted(KINDS))
def test_every_header_we_write_is_a_header_we_understand(kind):
    """Catches a typo in `template_headers` at its source.

    The round-trip above catches this too, but reports it as "unknown header"
    from inside a workbook. This names the offending string.
    """
    contract, _, _ = KINDS[kind]
    from app.ingest.sources.excel.reader import normalize_header

    for header in contract.template_headers:
        assert normalize_header(header) in contract.columns, header


def test_the_schedule_template_offers_every_column_we_read():
    """A contract column with no template column is data we can never receive."""
    missing = SCHEDULE_CONTRACT.canonical_fields() - SCHEDULE_CONTRACT.template_fields()

    assert missing == set(), f"schedule fields with no column: {sorted(missing)}"


def test_the_worklog_template_omits_exactly_one_optional_column():
    """Named rather than left to drift.

    `blocked_by` is read if present and not offered in the blank template: it
    names another row, and a free-text reference to a row is the kind of column
    the dependency resolver exists to replace.

    `log_date` used to be omitted too, which is what made an effort burn chart
    impossible - hours with no date are a total, not a series. It is now a
    column, and this assertion is the record that dropping it again is a
    decision rather than an accident.
    """
    missing = WORKLOG_CONTRACT.canonical_fields() - WORKLOG_CONTRACT.template_fields()

    assert missing == {"blocked_by"}


def test_the_demo_generator_and_the_template_describe_one_sheet():
    """The drift this refactor removed, kept removed.

    `gen_demo_data` writes positional row literals, so a header list that
    diverged from the contract would shift every value one column left - a
    corruption that still opens cleanly in Excel.
    """
    from scripts.gen_demo_data import SCHEDULE_HEADERS, WORKLOG_HEADERS

    assert SCHEDULE_HEADERS == list(SCHEDULE_CONTRACT.template_headers)
    assert WORKLOG_HEADERS == list(WORKLOG_CONTRACT.template_headers)


def test_the_notes_tab_is_not_the_sheet_the_reader_opens():
    """Guidance must be unreachable from the pipeline.

    It is prose in cells; if it shared a sheet with the table it would arrive as
    rejected rows on every sync.
    """
    workbook = template_workbook("schedule")

    assert "Notes" in workbook.sheetnames
    assert workbook.sheetnames[0] == "Activities"


def test_an_unknown_template_is_refused():
    with pytest.raises(ValueError) as caught:
        template_bytes("gantt")

    assert "schedule" in str(caught.value)


# --------------------------------------------------------------------------
# The status report
# --------------------------------------------------------------------------

def _importable(name: str) -> bool:
    """Whether a module can actually be imported, not merely located.

    `find_spec` is not enough - it answers "is there a file", which is true for
    a half-installed package that raises on import. Only `ModuleNotFoundError`
    is caught, matching `pytest.importorskip`: a genuinely absent extra should
    skip, and a *broken* install should fail loudly rather than hide.
    """
    try:
        importlib.import_module(name)
        return True
    except ModuleNotFoundError:
        return False


#: Skips only the report tests, not the whole module. A module-level
#: `importorskip` would take the template round-trips down with it, and those
#: need nothing but openpyxl - a core dependency, not an extra.
requires_docx = pytest.mark.skipif(
    not _importable("docx"), reason="python-docx is an optional extra"
)


def _bundle(**kw) -> InsightBundle:
    base = dict(
        project_id="excel:Project:1:HRMS",
        generated_at=datetime(2026, 3, 22, tzinfo=timezone.utc),
        as_of=datetime(2026, 3, 22, tzinfo=timezone.utc),
        narrative=(
            "What is at risk\nThe plan cannot hold.\n\n"
            "Why it is happening\nWBS-108 slipped.\n\n"
            "What it will impact\nDownstream tasks move.\n\n"
            "What to do next\nRe-baseline."
        ),
        findings=[
            Finding(
                id="schedule_inconsistent_major",
                category="schedule_risk",
                severity="high",
                headline="The plan cannot hold: 3 task(s) are late by up to 34 days.",
                recommendation="Re-baseline the affected tasks.",
                facts={"tasks_inconsistent": "3", "max_propagated_days": "34"},
                rule_trace=RuleTrace(
                    rule_id="schedule_inconsistent_major",
                    conditions=["max_propagated_days (34) >= 5"],
                    rationale="Five days is a working week.",
                ),
                evidence=[
                    EvidenceRef(raw_data_id=41, raw_table="_raw_excel_rows",
                                remark="Activities!row5")
                ],
            )
        ],
        data_quality=DataQuality(
            rows_rejected=8,
            changes_total=27,
            baseline_coverage=0.4,
            edges_stated=3,
            edges_inferred=1,
            depends_on_inferred_edges=True,
        ),
    )
    base.update(kw)
    return InsightBundle(**base)


def _explain() -> ExplainBundle:
    return ExplainBundle(
        project_id="excel:Project:1:HRMS",
        steps=[
            ForwardStep(
                entity_id="excel:Task:1:WBS-114",
                label="WBS-114",
                planned_end=date(2026, 4, 1),
                projected_end=date(2026, 4, 12),
                propagated_days=11,
            ),
            ForwardStep(
                entity_id="excel:Task:1:WBS-200",
                label="WBS-200",
                planned_end=date(2026, 5, 1),
            ),
        ],
        driving_path=["excel:Task:1:WBS-114"],
        project_slip_days=22,
        depends_on_inferred_edges=True,
    )


def _text_of(document) -> str:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        parts += [cell.text for row in table.rows for cell in row.cells]
    return "\n".join(parts)


@requires_docx
def test_the_report_states_the_findings_verbatim():
    """Rendered, never re-worded. The headline is already substituted prose."""
    from app.exports.report import build_report

    text = _text_of(build_report(_bundle()))

    assert "The plan cannot hold: 3 task(s) are late by up to 34 days." in text
    assert "Re-baseline the affected tasks." in text


@requires_docx
def test_the_report_carries_the_rule_trace_and_the_source_row():
    """Both are on the never-cut list, so both are asserted."""
    from app.exports.report import build_report

    text = _text_of(build_report(_bundle()))

    assert "max_propagated_days (34) >= 5" in text
    assert "Five days is a working week." in text
    assert "Activities!row5" in text
    assert "41" in text


@requires_docx
def test_the_report_never_hides_that_a_claim_rests_on_an_inference():
    """The document this product argues against is the one that omits this."""
    from app.exports.report import build_report

    text = _text_of(build_report(_bundle(), explain=_explain()))

    assert "inferred from the" in text
    assert "derived claim" in text


@requires_docx
def test_the_report_formats_a_percentage_the_way_the_assembler_does():
    """Invariant 1 across a .docx.

    `baseline_coverage` is 0.4 in the bundle. If this module formatted it, the
    report would say "0.4" or "40.0%"; `format_fact` says "40%", which is what
    the screen says.
    """
    from app.exports.report import build_report

    text = _text_of(build_report(_bundle()))

    assert "40%" in text
    assert "0.4" not in text


@requires_docx
def test_the_projection_table_puts_the_driving_path_first():
    """It is the only part of the table anyone acts on."""
    from app.exports.report import build_report

    table = build_report(_bundle(), explain=_explain()).tables[0]
    first_task = table.rows[1].cells[0].text

    assert first_task == "WBS-114"
    assert table.rows[1].cells[4].text == "yes"


@requires_docx
def test_the_report_says_which_wrote_the_summary():
    """A reader who cannot tell a model from a template cannot judge either."""
    from app.exports.report import build_report

    template_text = _text_of(build_report(_bundle()))
    model_text = _text_of(build_report(_bundle(narration_source="model")))

    assert "deterministic template" in template_text
    assert "language model" in model_text


@requires_docx
def test_a_bundle_with_no_findings_still_produces_a_report():
    from app.exports.report import build_report

    text = _text_of(build_report(_bundle(findings=[])))

    assert "Nothing in this project breaches a delivery threshold." in text
    # The caveat section is never dropped, even when there is nothing to report.
    assert "State changes observed" in text


@requires_docx
def test_report_bytes_is_a_real_docx():
    """Opens as a document, not merely as a zip of the right size."""
    from app.exports.report import report_bytes

    content = report_bytes(_bundle(), explain=_explain())

    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "r.docx"
        path.write_bytes(content)

        import docx

        assert "The plan cannot hold" in _text_of(docx.Document(str(path)))


# --------------------------------------------------------------------------
# The download routes
#
# Asserting `content-type`, not just the status. A Cloudflare Pages misconfig
# once served HTML with a 200 for every missing asset here, so "it returned
# 200" is not evidence that a file arrived - see CLAUDE.md section 6.
# --------------------------------------------------------------------------


@pytest.mark.parametrize("kind", sorted(KINDS))
def test_the_template_route_serves_a_spreadsheet(kind):
    from fastapi.testclient import TestClient

    from app.api.main import XLSX_TYPE, app

    response = TestClient(app).get(f"/api/template/{kind}.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_TYPE
    assert f"projectpulse_{kind}.xlsx" in response.headers["content-disposition"]
    # A real workbook starts with the zip magic; HTML served by mistake does not.
    assert response.content[:2] == b"PK"


def test_an_unknown_template_route_is_a_404_not_a_500():
    from fastapi.testclient import TestClient

    from app.api.main import app

    response = TestClient(app).get("/api/template/gantt.xlsx")

    assert response.status_code == 404
    assert "schedule" in response.json()["detail"]
