"""Uploading a document for a project through the browser.

The alternative ingestion path added alongside the project switcher: a
deployed host has no OneDrive-synced folder for `scripts.sync` to read, so
`POST /api/sources/upload` is how it gets data at all. These tests exercise
the whole thing - save the file, register it, run the same sync the PM's
button runs - against in-memory SQLite, no Docker.
"""

from __future__ import annotations

from datetime import datetime
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app import scope
from app.ingest.sources.excel import source as excel_source


@pytest.fixture()
def client():
    return TestClient(__import__("app.api.main", fromlist=["app"]).app)


@pytest.fixture(autouse=True)
def _isolated_registries(tmp_path, monkeypatch):
    """Each test gets its own data root and an empty upload registry.

    Two halves, because the state moved:

    **The data root** is still a directory, and still redirected at a temp
    one, because the demo's *generated* sheets live there and a test must not
    read or write the developer's real ones. `settings` is a frozen dataclass,
    so this replaces the module-level singleton rather than one of its fields
    - and on every module that imported it by name at its own top level
    (`app.api.main`, `app.ingest.sources.excel.source`), since a top-level
    `from x import y` binds a snapshot a patch to `x.y` alone would not reach.

    **The registries** are database rows now (see `app/models/uploads.py`),
    and the suite shares one SQLite file - so isolation means truncating two
    tables rather than handing out a temp directory. Without it, one test's
    imported project is in `scope.all_projects()` for every test that runs
    after it, and the portfolio, the picker and the risk form all see it.
    """
    import dataclasses

    from sqlalchemy import delete

    import app.api.main as main_module
    import app.config as config_module
    import app.ingest.sources.excel.source as excel_source_module
    from app.db import session_scope
    from app.models.uploads import RegisteredProject, UploadedSheet

    patched = dataclasses.replace(
        config_module.settings,
        state_dir=tmp_path / "state",
        data_root=tmp_path / "data",
    )
    monkeypatch.setattr(config_module, "settings", patched)
    monkeypatch.setattr(main_module, "settings", patched)
    monkeypatch.setattr(excel_source_module, "settings", patched)

    def clear():
        with session_scope() as session:
            session.execute(delete(UploadedSheet))
            session.execute(delete(RegisteredProject))

    clear()
    yield
    clear()


def _schedule_workbook(path: Path) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Activities"
    sheet.append(
        [
            "Task ID",
            "Activity",
            "Phase",
            "Milestone",
            "Status",
            "Owner",
            "Start",
            "Baseline Finish",
            "Planned Finish",
            "Progress",
            "Predecessor",
        ]
    )
    sheet.append(
        [
            "WBS-1",
            "Kickoff",
            "Planning",
            "No",
            "Done",
            "A. Owner",
            "2026-01-05",
            "2026-01-12",
            "2026-01-12",
            100,
            None,
        ]
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path.read_bytes()


def test_uploading_a_new_project_registers_and_ingests_it(client, tmp_path):
    contents = _schedule_workbook(tmp_path / "upload.xlsx")

    response = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "New Client Rollout"},
        files={"file": ("upload.xlsx", contents, "application/vnd.ms-excel")},
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ok"] is True
    assert body["project_id"].startswith("excel:Project:upload:")
    assert body["project_name"] == "New Client Rollout"
    assert body["rows_ok"] >= 1

    # Registered for real, not just echoed back in the response.
    found = scope.find(body["project_id"])
    assert found is not None
    assert found.name == "New Client Rollout"

    portfolio = client.get("/api/portfolio").json()
    ids = [row["project_id"] for row in portfolio["projects"]]
    assert body["project_id"] in ids


def test_a_second_upload_for_the_same_project_updates_rather_than_duplicates(client, tmp_path):
    contents = _schedule_workbook(tmp_path / "upload.xlsx")
    first = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "Repeat Co"},
        files={"file": ("upload.xlsx", contents, "application/vnd.ms-excel")},
    ).json()

    second = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_id": first["project_id"]},
        files={"file": ("upload.xlsx", contents, "application/vnd.ms-excel")},
    )

    assert second.status_code == 200, second.text
    body = second.json()
    assert body["ok"] is True
    assert body["project_id"] == first["project_id"]
    assert body["file_name"] == first["file_name"]

    watched = [w.file_name for w in excel_source.all_watched() if w.project_id == first["project_id"]]
    assert watched == [first["file_name"]]  # one entry, not two


def test_an_unknown_sheet_kind_is_refused(client, tmp_path):
    contents = _schedule_workbook(tmp_path / "upload.xlsx")

    response = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "not-a-real-kind", "project_name": "Whatever"},
        files={"file": ("upload.xlsx", contents, "application/vnd.ms-excel")},
    )

    assert response.status_code == 400
    assert "sheet_kind" in response.json()["detail"]


def test_a_new_project_needs_a_name(client, tmp_path):
    contents = _schedule_workbook(tmp_path / "upload.xlsx")

    response = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule"},
        files={"file": ("upload.xlsx", contents, "application/vnd.ms-excel")},
    )

    assert response.status_code == 400
    assert "project_name" in response.json()["detail"]


def test_an_unknown_project_id_is_refused(client, tmp_path):
    contents = _schedule_workbook(tmp_path / "upload.xlsx")

    response = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_id": "excel:Project:upload:nope"},
        files={"file": ("upload.xlsx", contents, "application/vnd.ms-excel")},
    )

    assert response.status_code == 400
    assert "unknown project" in response.json()["detail"]


def test_a_non_excel_file_is_refused_by_extension_before_anything_is_read(client):
    response = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "Wrong Format"},
        files={"file": ("notes.txt", b"hello", "text/plain")},
    )

    assert response.status_code == 400
    assert "Excel workbook" in response.json()["detail"]


# ==========================================================================
# Importing a document somebody already had
#
# Every test above uploads a workbook whose tab is called `Activities` - the
# name our own generated template uses, so it passes by construction. A PM
# importing a real document is uploading a file whose tab is called whatever
# Excel called it, and that case ingested zero rows while answering 200.
# ==========================================================================


def _schedule_workbook_named(path: Path, tab: str, *, task_id: str = "WBS-1") -> bytes:
    """The same workbook as `_schedule_workbook`, on a differently-named tab."""
    from openpyxl import load_workbook

    _schedule_workbook(path)
    workbook = load_workbook(path)
    workbook["Activities"].title = tab
    workbook[tab].cell(row=2, column=1).value = task_id
    workbook.save(path)
    return path.read_bytes()


@pytest.mark.parametrize("tab", ["Sheet1", "Schedule", "WBS", "Plan 2026"])
def test_a_document_is_imported_whatever_its_tab_is_called(client, tmp_path, tab):
    """The header *row* was always searched for rather than assumed; the tab
    name was the matching rigidity nobody had removed."""
    contents = _schedule_workbook_named(tmp_path / "real.xlsx", tab)

    response = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": f"Doc On {tab}"},
        files={"file": ("real.xlsx", contents, "application/vnd.ms-excel")},
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ok"] is True, body
    assert body["sheet_name"] == tab
    assert body["rows_ok"] >= 1


def test_the_conventional_tab_wins_when_a_workbook_has_several(client, tmp_path):
    """`preferred` is tried first, so nothing about an existing sheet moves."""
    from openpyxl import load_workbook

    path = tmp_path / "many.xlsx"
    _schedule_workbook(path)
    workbook = load_workbook(path)
    # A decoy that would also qualify, placed first so sheet order cannot be
    # what makes this pass.
    decoy = workbook.create_sheet("Old Plan", 0)
    for row in workbook["Activities"].iter_rows(values_only=True):
        decoy.append(row)
    workbook.save(path)

    response = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "Several Tabs"},
        files={"file": ("many.xlsx", path.read_bytes(), "application/vnd.ms-excel")},
    )

    assert response.json()["sheet_name"] == "Activities"


def test_the_resolved_tab_is_remembered_so_a_re_import_keeps_its_baseline(
    client, tmp_path
):
    """The sheet name is half the scope key.

    Re-deriving it per scan would be fine while a tab keeps its name and
    catastrophic when one is renamed: the scope changes, the differ finds no
    baseline, and every row comes back as a change with no lower bound. So it
    is resolved once and stored with the registration.
    """
    from openpyxl import load_workbook

    path = tmp_path / "real.xlsx"
    contents = _schedule_workbook_named(path, "Sheet1")
    first = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "Remembered"},
        files={"file": ("real.xlsx", contents, "application/vnd.ms-excel")},
    ).json()

    watched = {w.file_name: w for w in excel_source.all_watched()}
    assert watched[first["file_name"]].sheet_name == "Sheet1"

    # Re-import the same document with one date moved. A kept baseline means
    # the differ reports one change; a lost one reports every row as new.
    workbook = load_workbook(path)
    workbook["Sheet1"].cell(row=2, column=9).value = "2026-02-20"  # Planned Finish
    workbook.save(path)

    second = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_id": first["project_id"]},
        files={"file": ("real.xlsx", path.read_bytes(), "application/vnd.ms-excel")},
    ).json()

    assert second["ok"] is True, second
    assert second["changes_emitted"] == 1, second


def test_a_workbook_with_no_schedule_in_it_is_refused_and_says_what_is_missing(
    client, tmp_path
):
    """Refused, not registered. A project that ingests nothing is worse than
    no project: it sits on the portfolio as `no_data` and looks like a bug."""
    workbook = Workbook()
    workbook.active.title = "Notes"
    workbook.active.append(["Meeting notes"])
    workbook.active.append(["Talked about the budget"])
    path = tmp_path / "notes.xlsx"
    workbook.save(path)

    response = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "Just Notes"},
        files={"file": ("notes.xlsx", path.read_bytes(), "application/vnd.ms-excel")},
    )

    assert response.status_code == 400
    detail = response.json()["detail"]
    # Name the tabs we looked at and the column we need, or the person is sent
    # back to a file with nothing to change.
    assert "Notes" in detail
    assert "Task ID" in detail
    assert scope.find("excel:Project:upload:just-notes") is None


def _stored_bytes(file_name: str) -> bytes | None:
    from app.db import session_scope
    from app.models.uploads import UploadedSheet

    with session_scope() as session:
        row = session.get(UploadedSheet, file_name)
        return row.content if row else None


def test_an_imported_workbook_is_stored_where_a_deploy_cannot_lose_it(
    client, tmp_path
):
    """The bytes go in the database, beside the registration.

    They used to be written to `settings.data_root`, which is a container's
    local disk on a deployed host - so a document imported through the browser
    was gone after the next release, while the sheet went on being watched
    with nothing behind it. This is what makes the website usable without a
    local copy of the app.
    """
    contents = _schedule_workbook_named(tmp_path / "real.xlsx", "Sheet1")

    body = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "Durable"},
        files={"file": ("Q2 plan.xlsx", contents, "application/vnd.ms-excel")},
    ).json()

    assert _stored_bytes(body["file_name"]) == contents
    # Nothing was written to the folder - the whole point.
    assert not (Path(config_module().data_root) / body["file_name"]).exists()


def test_a_stored_workbook_is_ingested_with_no_file_on_disk(client, tmp_path):
    """What a machine looks like after a deploy: the row is there, the disk is
    empty, and the sheet still syncs."""
    contents = _schedule_workbook_named(tmp_path / "real.xlsx", "Sheet1")
    first = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "Survives"},
        files={"file": ("plan.xlsx", contents, "application/vnd.ms-excel")},
    ).json()
    assert first["ok"] is True

    # A fresh container: nothing on disk, same database.
    data_root = Path(config_module().data_root)
    for stale in data_root.glob("*.xlsx"):
        stale.unlink()

    from app.ingest.sources.excel.source import all_watched
    from app.ingest.sources.excel.transport import StoredSheetSource

    watched = {w.file_name: w for w in all_watched()}[first["file_name"]]
    fetched = StoredSheetSource().fetch(watched)

    assert fetched is not None, "the stored workbook was not found"
    try:
        assert fetched.local_path.read_bytes() == contents
        # A human has to be told where this came from, and there is no path to
        # give them - so it names the file they uploaded.
        assert fetched.display_uri == "upload://plan.xlsx"
    finally:
        fetched.release()
    # The temp copy is cleaned up, or a long-running server fills its disk.
    assert not fetched.local_path.exists()


def test_a_refused_workbook_does_not_replace_the_one_already_syncing(
    client, tmp_path
):
    """The workbook is inspected from a temp file and only stored once it has
    been accepted, so a bad re-upload cannot displace the copy the differ is
    diffing against."""
    good = _schedule_workbook_named(tmp_path / "good.xlsx", "Sheet1")
    first = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "Keep Mine"},
        files={"file": ("good.xlsx", good, "application/vnd.ms-excel")},
    ).json()

    junk = Workbook()
    junk.active.title = "Nope"
    junk.active.append(["not a schedule"])
    junk_path = tmp_path / "junk.xlsx"
    junk.save(junk_path)

    refused = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_id": first["project_id"]},
        files={"file": ("junk.xlsx", junk_path.read_bytes(), "application/vnd.ms-excel")},
    )

    assert refused.status_code == 400
    assert _stored_bytes(first["file_name"]) == good


def test_two_documents_that_number_their_tasks_the_same_way_do_not_collide(
    client, tmp_path
):
    """The one that could lose a project's data rather than just fail.

    Every Excel project shares one `connection_id`, and a row key is unique
    only within its own sheet - so while `domain_id` namespaced a Task by
    connection + row key alone, importing team B's plan silently took team
    A's tasks. The demo generators prefix their ids per project to dodge it;
    a document somebody uploads cannot be asked to.
    """
    from sqlalchemy import func, select

    from app.db import session_scope
    from app.models.domain import Task

    a = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "Team A"},
        files={
            "file": (
                "a.xlsx",
                _schedule_workbook_named(tmp_path / "a.xlsx", "Sheet1", task_id="1"),
                "application/vnd.ms-excel",
            )
        },
    ).json()
    b = client.post(
        "/api/sources/upload",
        data={"sheet_kind": "schedule", "project_name": "Team B"},
        files={
            "file": (
                "b.xlsx",
                _schedule_workbook_named(tmp_path / "b.xlsx", "WBS", task_id="1"),
                "application/vnd.ms-excel",
            )
        },
    ).json()

    assert a["project_id"] != b["project_id"]
    with session_scope() as session:
        for project_id in (a["project_id"], b["project_id"]):
            count = session.scalar(
                select(func.count())
                .select_from(Task)
                .where(Task.project_id == project_id)
            )
            assert count == 1, f"{project_id} has {count} task(s), expected 1"


def test_a_missing_tab_is_reported_as_a_missing_tab(tmp_path):
    """Not as "sheet has no 'task_id' column", which is what it used to say.

    `read_sheet` returns its own rejection naming the tabs it found, and
    `resolve_identities` used to raise `MissingKeyColumn` straight over the
    top of it - because there were no rows, not because the column was
    absent. The person was then told to add a Task ID column to a sheet whose
    Task ID column was right there, and the real cause was never printed.
    """
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    from app.ingest.sources.excel.reader import SCHEDULE_CONTRACT
    from app.ingest.sources.excel.ingest import ingest_sheet
    from app.models.base import Base

    path = tmp_path / "renamed.xlsx"
    _schedule_workbook_named(path, "Sheet1")

    engine = create_engine("sqlite://", future=True)
    Base.metadata.create_all(engine)
    with sessionmaker(bind=engine, future=True)() as session:
        report = ingest_sheet(
            session,
            source="excel",
            connection_id=1,
            file_path=path,
            sheet_name="Activities",  # what a caller expected, not what is there
            contract=SCHEDULE_CONTRACT,
            project_id="excel:Project:1:X",
            now=__import__("datetime").datetime(2026, 3, 2, 9, 0),
            sync_run_id=1,
            logical_name="renamed.xlsx",
        )

    assert not report.ok
    assert "not found" in report.error
    assert "Sheet1" in report.error  # the tabs the workbook does have
    assert "task_id" not in report.error


def config_module():
    """The patched `settings` the isolation fixture installed."""
    import app.config

    return app.config.settings

# --------------------------------------------------------------------------
# Removing a project. The one operation with nothing to rebuild from.
# --------------------------------------------------------------------------


def test_a_built_in_seed_project_is_refused_with_a_reason():
    """`app/scope.py` declares HRMS in code, not in the registry. Deleting its
    data would empty it without removing it - leaving a project on the portfolio
    with nothing in it, which looks exactly like one awaiting its first sync and
    is a worse state than the one somebody was trying to leave."""
    from app.db import session_scope
    from app.projects import plan, remove

    with session_scope() as session:
        outcome = plan(session, "excel:Project:1:HRMS")
        assert outcome.removable is False
        assert "scope.py" in outcome.reason

        with pytest.raises(ValueError):
            remove(session, "excel:Project:1:HRMS")


def test_an_unknown_project_is_refused_rather_than_silently_doing_nothing():
    from app.db import session_scope
    from app.projects import plan

    with session_scope() as session:
        outcome = plan(session, "nothing:Project:9:NOPE")
        assert outcome.removable is False
        assert "no project" in outcome.reason


def test_the_plan_counts_what_would_go_without_removing_any_of_it():
    """The preview has to be free of side effects, or pressing Remove and then
    cancelling would already have done half the job."""
    from sqlalchemy import func, select

    from app.db import session_scope
    from app.models.domain import Task
    from app.projects import plan

    with session_scope() as session:
        before = session.scalar(select(func.count()).select_from(Task)) or 0
        plan(session, "excel:Project:1:HRMS")
        plan(session, "nothing:Project:9:NOPE")
        assert (session.scalar(select(func.count()).select_from(Task)) or 0) == before


def test_a_removal_plan_names_the_rows_no_sync_can_rebuild():
    """A risk somebody typed and a board somebody arranged have no source system
    behind them. A confirmation that does not lead with those is asking the
    wrong question."""
    from app.projects import RemovalPlan

    outcome = RemovalPlan(
        project_id="p", name="P", removable=True,
        counts={"tasks": 10, "risks": 2, "dashboards": 1, "qa_items": 0},
    )
    assert outcome.irreplaceable == {"risks": 2, "dashboards": 1}

def test_one_jira_export_becomes_both_sheets_when_it_carries_effort(tmp_path):
    """The form asks for a file, not for which half of it to read.

    An issue row holds a plan and a record of effort, and this app keeps those
    in separate contracts - but that is an internal arrangement, and making
    somebody upload the same file twice leaks it onto the form.
    """
    from datetime import datetime

    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import NotAJiraExport, convert_workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["Filter"])
    sheet.append(["Displaying 1 issues at 13/Sep/26 9:00 AM."])
    sheet.append(["Project", "Key", "Summary", "Status", "Assignee", "Due Date",
                  "Original Estimate", "Time Spent"])
    sheet.append(["P", "P-1", "Design", "To Do", "Ann", datetime(2026, 9, 30), 8, 12])
    path = tmp_path / "with_effort.xlsx"
    book.save(path)

    schedule, _ = convert_workbook(path, kind="schedule")
    worklog, _ = convert_workbook(path, kind="worklog")
    assert schedule and worklog


def test_a_jira_export_without_effort_still_imports_its_schedule(tmp_path):
    """The common case, and the one that must not regress: plenty of Jira
    projects never fill in an estimate, and refusing the whole upload for that
    would reject a perfectly good schedule."""
    from datetime import datetime

    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import NotAJiraExport, convert_workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["Filter"])
    sheet.append(["Displaying 1 issues at 13/Sep/26 9:00 AM."])
    sheet.append(["Project", "Key", "Summary", "Status", "Assignee", "Due Date"])
    sheet.append(["P", "P-1", "Design", "To Do", "Ann", datetime(2026, 9, 30)])
    path = tmp_path / "no_effort.xlsx"
    book.save(path)

    schedule, _ = convert_workbook(path, kind="schedule")
    assert schedule

    # The worklog half refuses - and the route treats that as "no effort here",
    # not as a failed import.
    with pytest.raises(NotAJiraExport):
        convert_workbook(path, kind="worklog")


def test_the_upload_form_offers_one_jira_option_not_two():
    """Pinned because the two-option version shipped and was wrong: a person has
    one file and should not be asked which half of it to read."""
    from pathlib import Path

    import app.api.main as main

    html = (Path(main.__file__).parent / "static" / "settings.html").read_text(
        encoding="utf-8"
    )
    assert html.count('value="jira_export"') == 1
    assert "jira_worklog" not in html


def test_a_backlog_appended_below_the_export_is_read_rather_than_dropped(tmp_path):
    """The case this was written for: a sheet that is a Jira export *plus* rows
    somebody pasted underneath it.

    Jira numbers every issue it writes, so "one issue is one row with a Key"
    reads an export correctly - and silently discards a backlog appended below
    one, which is how a 191-row sheet arrived as 17 tasks with nothing finished.
    """
    from datetime import datetime

    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import read_export

    book = Workbook()
    sheet = book.active
    sheet.append(["Filter"])
    sheet.append(["Displaying 1 issues at 13/Sep/26 9:00 AM."])
    sheet.append(["Project", "Key", "Summary", "Status", "Assignee", "Due Date",
                  "Description"])
    sheet.append(["P", "P-1", "Design", "To Do", "Ann", datetime(2026, 9, 30), "a"])
    # The export's own wrapped rich text: the Description column and nothing else.
    sheet.append([None, None, None, None, None, None, "continued on this row"])
    # Appended by hand: no Key, but a Summary and a Status, which no wrapping
    # produces.
    sheet.append(["P", None, "Ship the thing", "Release it", None, None, None])
    path = tmp_path / "appended.xlsx"
    book.save(path)

    _, rows = read_export(path)
    assert len(rows) == 2
    assert [r[2] for r in rows] == ["Design", "Ship the thing"]


def test_an_appended_row_gets_a_derived_id_that_cannot_pass_for_a_jira_key(tmp_path):
    """`Task ID` is the schedule contract's key field, so a keyless row needs
    one - but it must not look like an issue somebody can open, and it has to
    survive the sheet being re-sorted or the differ reports every task as
    deleted and re-added."""
    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import (
        SYNTHETIC_PREFIX,
        convert,
        read_export,
    )

    def _sheet(order):
        book = Workbook()
        page = book.active
        page.append(["Filter"])
        page.append(["Displaying 0 issues at 13/Sep/26 9:00 AM."])
        page.append(["Project", "Key", "Summary", "Status"])
        for summary in order:
            page.append(["P", None, summary, "Release it"])
        path = tmp_path / f"{'_'.join(order)}.xlsx"
        book.save(path)
        return path

    def _ids(order):
        headers, rows = read_export(_sheet(order))
        records, _ = convert(headers, rows)
        return {r["Activity"]: r["Task ID"] for r in records}

    ids = _ids(["Ship it", "Test it"])
    assert all(v.startswith(SYNTHETIC_PREFIX) for v in ids.values())
    # Not shaped like PROJ-12, so nothing downstream treats it as a real issue.
    assert not any(re.fullmatch(r"[A-Z][A-Z0-9_]*-\d+", v) for v in ids.values())
    # Hashed from the summary, so re-ordering the sheet does not rename anything.
    assert _ids(["Test it", "Ship it"]) == ids


def test_a_derived_id_never_collides_with_a_real_key_further_up_the_sheet(tmp_path):
    """The keyed issues sit *above* the appended ones, so a check against only
    what has been seen so far would be looking the wrong way."""
    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import convert, read_export

    book = Workbook()
    sheet = book.active
    sheet.append(["Filter"])
    sheet.append(["Displaying 2 issues at 13/Sep/26 9:00 AM."])
    sheet.append(["Project", "Key", "Summary", "Status"])
    sheet.append(["P", "P-1", "Design", "To Do"])
    # Two appended rows sharing a summary: two items on the board, not one.
    sheet.append(["P", None, "Ship it", "Release it"])
    sheet.append(["P", None, "Ship it", "Release it"])
    path = tmp_path / "collide.xlsx"
    book.save(path)

    headers, rows = read_export(path)
    records, _ = convert(headers, rows)
    ids = [r["Task ID"] for r in records]
    assert len(ids) == len(set(ids)) == 3


def test_the_coverage_report_says_when_it_derived_the_ids(tmp_path):
    """Reading rows Jira did not write changes what every count below is
    counting, so it is named rather than left for somebody to notice."""
    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import convert_workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["Filter"])
    sheet.append(["Displaying 1 issues at 13/Sep/26 9:00 AM."])
    sheet.append(["Project", "Key", "Summary", "Status"])
    sheet.append(["P", "P-1", "Design", "To Do"])
    sheet.append(["P", None, "Ship it", "Release it"])
    path = tmp_path / "reported.xlsx"
    book.save(path)

    _, cover = convert_workbook(path, kind="schedule")
    assert any("carry no Jira Key" in n for n in cover["notes"])


def test_a_shipped_status_counts_as_done_rather_than_as_open_work():
    """`Release it` is the terminal state in the Jira workflow this app reads.
    Unmapped it became OTHER, which every downstream count reads as open - so a
    released board reported nothing finished and suppressed four findings that
    are gated on exactly that."""
    from app.ingest.sources.excel.convertor import _normalize_status

    assert _normalize_status("Release it") == "DONE"
    assert _normalize_status("Released") == "DONE"
    # Unchanged: an ambiguous status is still an honest unknown.
    assert _normalize_status("Ready for release") == "OTHER"


def test_a_row_with_no_parent_is_banded_by_its_component(tmp_path):
    """A grouping column that groups nothing is not doing its job: an appended
    backlog row carries no parent issue and cannot, but it does say which
    feature area it belongs to."""
    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import convert, read_export

    book = Workbook()
    sheet = book.active
    sheet.append(["Filter"])
    sheet.append(["Displaying 1 issues at 13/Sep/26 9:00 AM."])
    sheet.append(["Project", "Key", "Summary", "Status", "Product", "Component/s"])
    # A real parent still wins, and the component beside it is not consulted.
    sheet.append(["P", "P-1", "Design", "To Do", "Platform [P-9]", "UI/UX"])
    sheet.append(["P", None, "Ship it", "Release it", None, "Cowork Chat"])
    path = tmp_path / "banded.xlsx"
    book.save(path)

    headers, rows = read_export(path)
    records, chosen = convert(headers, rows)
    assert chosen["Milestone"] == "Product"
    assert [r["Milestone"] for r in records] == ["Platform", "Cowork Chat"]


def test_the_fallback_never_bands_a_row_by_a_url(tmp_path):
    """Same refusal the chosen source makes: a link to a thing is not the name
    of one."""
    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import convert, read_export

    book = Workbook()
    sheet = book.active
    sheet.append(["Filter"])
    sheet.append(["Displaying 0 issues at 13/Sep/26 9:00 AM."])
    sheet.append(["Project", "Key", "Summary", "Status", "Component/s"])
    sheet.append(["P", None, "Ship it", "Release it", "https://wiki/components"])
    path = tmp_path / "url_band.xlsx"
    book.save(path)

    headers, rows = read_export(path)
    records, _ = convert(headers, rows)
    assert records[0]["Milestone"] is None


def test_a_date_this_team_keeps_in_the_description_is_read(tmp_path):
    """This Jira defines forty date columns and populates none of them; the
    delivery rows carry `Ngay nhan:` inside the Description instead, as a fixed
    `Label: value` block the team maintains by hand.

    Reading it is the difference between four tasks drawn as bare dots and four
    drawn as overdue bars - the worst news in the export. Jira stores the field
    as a number, so the value arrives as a raw Excel serial.
    """
    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import convert, read_export

    book = Workbook()
    sheet = book.active
    sheet.append(["Filter"])
    sheet.append(["Displaying 0 issues at 13/Sep/26 9:00 AM."])
    sheet.append(["Project", "Key", "Summary", "Status", "Due Date", "Description"])
    sheet.append(["P", None, "Refactor code BE", "In Progress", datetime(2026, 8, 31),
                  "PO: HoachBV\nBA: FSG\nNgay nhan: 46246\nGhi chu: DuyLH19"])
    path = tmp_path / "desc_date.xlsx"
    book.save(path)

    headers, rows = read_export(path)
    records, _ = convert(headers, rows)
    assert records[0]["Start"].isoformat() == "2026-08-12"


def test_a_real_start_column_beats_the_description(tmp_path):
    """The block is a fallback for rows that have no column, never a source that
    competes with one."""
    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import convert, read_export

    book = Workbook()
    sheet = book.active
    sheet.append(["Filter"])
    sheet.append(["Displaying 0 issues at 13/Sep/26 9:00 AM."])
    sheet.append(["Project", "Key", "Summary", "Status", "Start date", "Description"])
    sheet.append(["P", None, "Refactor", "In Progress", datetime(2026, 9, 1),
                  "Ngay nhan: 46246"])
    path = tmp_path / "column_wins.xlsx"
    book.save(path)

    headers, rows = read_export(path)
    records, _ = convert(headers, rows)
    assert records[0]["Start"].isoformat() == "2026-09-01"


@pytest.mark.parametrize(
    "body",
    [
        # A number too small to be a date is a quantity. 7 as 1900-01-07 would be
        # a confident wrong date on a chart.
        "Ngay nhan: 7",
        # A year, not a day.
        "Ngay nhan: 2026",
        "Ngay nhan: not a date",
        # The words in passing, not as a labelled field.
        "We agreed the ngay nhan would be next week.",
        # A different label entirely.
        "Ghi chu: 46246",
    ],
)
def test_the_description_block_refuses_anything_that_is_not_a_date(tmp_path, body):
    """It reads one labelled field, not prose. Widening it to anything that
    parses would make every sentence in a Jira description a date source."""
    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import convert, read_export

    book = Workbook()
    sheet = book.active
    sheet.append(["Filter"])
    sheet.append(["Displaying 0 issues at 13/Sep/26 9:00 AM."])
    sheet.append(["Project", "Key", "Summary", "Status", "Description"])
    sheet.append(["P", None, "Refactor", "In Progress", body])
    path = tmp_path / f"refused_{abs(hash(body))}.xlsx"
    book.save(path)

    headers, rows = read_export(path)
    records, _ = convert(headers, rows)
    assert records[0]["Start"] is None
