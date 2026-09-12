"""Guards for the commands a judge actually types.

Every failure in this file has already happened once for real:

* `python -m scripts.replay` under the system Python died on a missing
  `sqlalchemy`, because the dependencies live in `.venv`;
* the re-exec that fixed that then broke `import scripts`, because after `execv`
  `sys.path[0]` is the script's directory rather than the repo root;
* `scripts/demo.py` chose its database *after* importing `app.db`, which reads
  `DATABASE_URL` at import - so the console silently talked to a Postgres nobody
  had started;
* the replay timeline lost one of its two no-edit syncs, which removes every
  causal chain while leaving the page looking populated.

None of those are caught by testing the library. They are caught here.
"""

from __future__ import annotations

import ast
import os
import subprocess
import sys
from pathlib import Path

import pytest

from scripts import _bootstrap
from scripts.replay import TIMELINE

REPO = Path(__file__).resolve().parent.parent
SCRIPTS = REPO / "scripts"

#: Entry points a human is documented to run.
ENTRY_POINTS = ("demo", "sync", "replay", "gen_demo_data", "gen_jira_data", "index_code")


# --------------------------------------------------------------------------
# Import order — the trap that made `scripts.demo` use the wrong database
# --------------------------------------------------------------------------


def _first_import_lines(path: Path) -> list[tuple[int, str]]:
    """Top-level imports in source order, as (line, dotted name)."""
    tree = ast.parse(path.read_text(encoding="utf-8"))
    found: list[tuple[int, str]] = []
    for node in tree.body:
        if isinstance(node, ast.Import):
            for alias in node.names:
                found.append((node.lineno, alias.name))
        elif isinstance(node, ast.ImportFrom) and node.module:
            found.append((node.lineno, node.module))
    return found


@pytest.mark.parametrize("name", ["demo", "sync"])
def test_bootstrap_is_imported_before_any_app_module(name):
    """`app.config` freezes DATABASE_URL when it is imported.

    Any database choice made after the first `app.*` import is dead code. This is
    why the import order in those files looks wrong - it is deliberate, and this
    test is what stops someone tidying it back.
    """
    imports = _first_import_lines(SCRIPTS / f"{name}.py")

    bootstrap_line = min(
        (line for line, mod in imports if mod.startswith("scripts._bootstrap")),
        default=None,
    )
    first_app_line = min(
        (line for line, mod in imports if mod == "app" or mod.startswith("app.")),
        default=None,
    )

    assert bootstrap_line is not None, f"{name}.py does not bootstrap"
    if first_app_line is not None:
        assert bootstrap_line < first_app_line


def test_demo_no_longer_sets_the_database_after_importing_app():
    """The original bug, pinned so it cannot come back."""
    source = (SCRIPTS / "demo.py").read_text(encoding="utf-8")

    assert 'os.environ.setdefault("DATABASE_URL"' not in source


# --------------------------------------------------------------------------
# _bootstrap
# --------------------------------------------------------------------------


def test_the_venv_interpreter_is_found():
    interpreter = _bootstrap._venv_python()

    assert interpreter is not None, "no .venv found; the re-exec cannot work"
    assert interpreter.exists()


def test_default_database_does_not_override_an_explicit_choice(monkeypatch):
    monkeypatch.setenv("DATABASE_URL", "sqlite:///chosen.db")

    assert _bootstrap.default_database() == "sqlite:///chosen.db"


def test_default_database_falls_back_to_sqlite(monkeypatch):
    """The CLI must run with nothing installed and no container."""
    monkeypatch.delenv("DATABASE_URL", raising=False)

    assert _bootstrap.default_database() == _bootstrap.DEFAULT_SQLITE
    assert "sqlite" in os.environ["DATABASE_URL"]


def test_ensure_venv_is_a_noop_once_guarded(monkeypatch):
    """The guard is what stops an exec loop if a venv lacks the dependencies."""
    monkeypatch.setenv(_bootstrap._GUARD, "1")
    called = []
    monkeypatch.setattr(os, "execv", lambda *a: called.append(a))

    _bootstrap.ensure_venv()

    assert called == []


def test_ensure_venv_does_not_relaunch_when_already_in_the_venv(monkeypatch):
    monkeypatch.delenv(_bootstrap._GUARD, raising=False)
    interpreter = _bootstrap._venv_python()
    monkeypatch.setattr(sys, "executable", str(interpreter))
    called = []
    monkeypatch.setattr(os, "execv", lambda *a: called.append(a))

    _bootstrap.ensure_venv()

    assert called == []


def test_a_relaunch_keeps_the_repo_importable(monkeypatch):
    """After execv, `sys.path[0]` is the script's directory.

    Without PYTHONPATH the relaunched process cannot import `scripts` or `app` -
    which is exactly how the first version of this failed.
    """
    monkeypatch.delenv(_bootstrap._GUARD, raising=False)
    monkeypatch.delenv("PYTHONPATH", raising=False)
    monkeypatch.setattr(sys, "executable", r"C:\some\other\python.exe")
    captured: list[tuple] = []
    monkeypatch.setattr(os, "execv", lambda *a: captured.append(a))

    _bootstrap.ensure_venv()

    assert captured, "expected a relaunch"
    assert str(REPO) in os.environ["PYTHONPATH"]
    # And it relaunches via `-m`, not by path, for the same reason.
    argv = captured[0][1]
    assert "-m" in argv


# --------------------------------------------------------------------------
# The replay timeline
# --------------------------------------------------------------------------


def test_the_timeline_keeps_its_no_edit_syncs():
    """Dropping one removes every causal chain, silently.

    The data would still look fine: the page fills with schedule findings and no
    "why", which is the worst failure mode available - plausible and wrong.
    """
    no_edit = [entry for entry in TIMELINE if entry[0] is None and entry[1] == "excel"]

    assert len(no_edit) == 2


def test_the_timeline_is_chronological():
    times = [when for _, _, when in TIMELINE]

    assert times == sorted(times)


def test_every_demo_step_appears_exactly_once_and_in_order():
    steps = [step for step, _, _ in TIMELINE if step is not None]

    assert steps == [0, 1, 2, 3]


def test_a_scan_falls_between_each_edit_and_the_next():
    """The property the whole precision model needs.

    Two bounded changes in adjacent windows share a boundary instant and can
    never be ordered, so every edit must be followed by a sync before the next
    edit lands.
    """
    edits = [i for i, (step, _, _) in enumerate(TIMELINE) if step is not None]
    for first, second in zip(edits, edits[1:]):
        between = TIMELINE[first + 1 : second]
        assert between, f"no sync between step {TIMELINE[first][0]} and the next"


# --------------------------------------------------------------------------
# The entry points actually start
# --------------------------------------------------------------------------


@pytest.mark.parametrize("name", ENTRY_POINTS)
def test_every_entry_point_is_importable(name):
    """A syntax error or bad import in a script is invisible until someone runs it."""
    module = __import__(f"scripts.{name}", fromlist=["main"])

    assert module is not None


@pytest.mark.parametrize("name", ["sync", "replay", "gen_demo_data", "index_code"])
def test_every_entry_point_offers_help_without_touching_a_database(name, tmp_path):
    """`--help` is the cheapest possible smoke test.

    It proves the module imports, its arguments parse, and nothing at import time
    needs a database - the failure that made every command hang for minutes.
    """
    env = dict(os.environ)
    env["DATABASE_URL"] = f"sqlite:///{(tmp_path / 'x.db').as_posix()}"
    env["PULSE_BOOTSTRAPPED"] = "1"  # already the right interpreter

    result = subprocess.run(
        [sys.executable, "-m", f"scripts.{name}", "--help"],
        cwd=REPO,
        env=env,
        capture_output=True,
        text=True,
        timeout=90,
    )

    assert result.returncode == 0, result.stderr
    assert "usage" in (result.stdout + result.stderr).lower()


def test_the_index_generator_produces_a_marker_and_real_entries():
    """WORKLOG.md's index is generated; a broken generator makes it stale silently."""
    from scripts.index_code import MARKER, render

    output = render()

    assert output.startswith(MARKER)
    assert "app/intelligence/assembler.py" in output
    assert "build_bundle(" in output


def test_the_index_reads_source_rather_than_importing_it():
    """It has to work mid-refactor, when the package does not import."""
    source = (SCRIPTS / "index_code.py").read_text(encoding="utf-8")

    assert "import ast" in source
    assert "importlib" not in source


def test_the_index_generator_can_report_staleness_without_writing():
    """`--check` exists so the index cannot rot unnoticed.

    It also means `--help` and unknown flags behave, rather than the generator
    ignoring its arguments and rewriting the file regardless - which is what it
    did before, so running `--help` had a side effect.
    """
    result = subprocess.run(
        [sys.executable, "-m", "scripts.index_code", "--check"],
        cwd=REPO,
        capture_output=True,
        text=True,
        timeout=90,
    )

    assert result.returncode in (0, 1)
    assert "index is" in (result.stdout + result.stderr)


# --------------------------------------------------------------------------
# Console encoding — a Windows console is not always UTF-8
# --------------------------------------------------------------------------

CLI_SOURCES = [p for p in SCRIPTS.rglob("*.py") if "__pycache__" not in str(p)] + [
    REPO / "app" / "intelligence" / "explain.py"
]


@pytest.mark.parametrize("path", CLI_SOURCES, ids=lambda p: p.name)
def test_no_cli_source_contains_a_character_the_console_cannot_print(path):
    """An em dash crashed `scripts.sync explain` outright.

    The development console here is cp932, and a judge's may be cp1252 or
    anything else. `argparse(description=__doc__)` also prints module docstrings
    as --help text, so a docstring is just as exposed as a print(). ASCII-only is
    the only rule that holds on every machine.
    """
    text = path.read_text(encoding="utf-8")
    offenders = sorted({c for c in text if ord(c) > 127})

    assert not offenders, f"{path.name} contains {offenders}"


@pytest.mark.parametrize("name", ["sync", "replay", "gen_demo_data", "index_code"])
def test_help_text_survives_a_legacy_codepage(name):
    """`--help` renders the module docstring; it must encode on cp932 and cp1252."""
    module = __import__(f"scripts.{name}", fromlist=["main"])

    for codec in ("cp932", "cp1252", "ascii"):
        (module.__doc__ or "").encode(codec)


# --------------------------------------------------------------------------
# explain — the arithmetic, shown
# --------------------------------------------------------------------------


def test_explain_renders_a_checkable_table_and_working():
    from datetime import date

    from app.intelligence import explain
    from app.intelligence.schedule.graph import EdgeRecord, TaskNode, build_graph
    from app.intelligence.schedule.impact import project_schedule

    tasks = [
        TaskNode("A", "p", title="Env", start_date=date(2026, 2, 16),
                 planned_end=date(2026, 3, 16), baseline_end=date(2026, 3, 4)),
        TaskNode("B", "p", title="Build", start_date=date(2026, 3, 5),
                 planned_end=date(2026, 4, 1), baseline_end=date(2026, 3, 20)),
    ]
    schedule = build_graph(tasks, [EdgeRecord("A", "B")])
    impact = project_schedule(schedule)

    table = explain.table(schedule, impact)
    working = explain.working(schedule, impact, "B")

    # Every column header present, and one row per task.
    for header, _ in explain.HEADERS:
        assert header in table
    assert table.count("\n") == 3  # header, rule, two rows

    # The terminal rendering is the same three-part shape as the page, because a
    # reader following along on paper needs the same structure either way.
    for section in ("INPUT", "ALGORITHM", "OUTPUT"):
        assert section in working
    for field in ("formula", "values", "result"):
        assert field in working


def test_explain_truncates_a_synthetic_key_rather_than_breaking_alignment():
    """A `~anon-<hex>` key is far wider than any real Task ID."""
    from datetime import date

    from app.intelligence import explain
    from app.intelligence.schedule.graph import TaskNode, build_graph
    from app.intelligence.schedule.impact import project_schedule

    long_key = "~anon-ef0576ffa2b2a0f2"
    tasks = [
        TaskNode(long_key, "p", start_date=date(2026, 3, 1),
                 planned_end=date(2026, 3, 5), baseline_end=date(2026, 3, 5))
    ]
    schedule = build_graph(tasks, [])

    lines = explain.table(schedule, project_schedule(schedule)).splitlines()

    assert len(lines[0]) == len(lines[2]), "row width does not match the header"
    assert long_key not in lines[2]


def test_explain_shows_a_dash_where_a_number_is_unknown():
    """Printing 0 for "no baseline" reads as "on time", which is a lie."""
    from datetime import date

    from app.intelligence import explain
    from app.intelligence.schedule.graph import TaskNode, build_graph
    from app.intelligence.schedule.impact import project_schedule

    tasks = [TaskNode("A", "p", start_date=date(2026, 3, 1),
                      planned_end=date(2026, 3, 5))]  # no baseline
    schedule = build_graph(tasks, [])

    impact = project_schedule(schedule)
    row = explain.rows(schedule, impact)[0]
    derived = explain.derive(schedule, impact, "A")

    # The overview shows a dash; the derivation says so in words.
    assert row.propagated == "0"
    assert any(o.value == explain.ABSENT for o in derived.inputs)
    assert any(o.value == "-" for o in derived.outputs)


# --------------------------------------------------------------------------
# publish — the static snapshot
# --------------------------------------------------------------------------


def test_the_snapshot_never_overwrites_the_architecture_page():
    """`site/index.html` is the one live URL people already have.

    A publisher that clobbered it would take down something that works in order
    to add something new.
    """
    from scripts.publish import APP_ROUTES, STATIC_PAGES, _published_pages

    assert "index.html" not in _published_pages()
    assert "index" not in APP_ROUTES
    assert "index.html" not in STATIC_PAGES


def test_the_snapshot_publishes_assets_at_the_url_the_pages_request():
    """The failure this replaces was invisible to every status check.

    The pages linked `/static/shell.css`; the publisher wrote `site/shell.css`.
    Cloudflare Pages answers an unmatched path with `index.html` and status
    **200**, so `curl -w %{http_code}` said 200 while the browser refused to
    apply HTML as a stylesheet - a whole site of unstyled pages, no chart, and
    white-on-white text where an inline rule set `color:#fff` against a token
    that no longer resolved.

    Destinations are now derived from the reference, so they cannot disagree.
    """
    from scripts.publish import _referenced_assets, _source_for

    referenced = _referenced_assets()
    assert referenced, "no assets found - has the app been built?"

    for url in referenced:
        assert url.startswith("/static/"), url
        assert _source_for(url).exists(), f"{url} -> {_source_for(url)} missing"


def test_the_snapshot_covers_the_hashed_bundle_paths():
    """Vite re-hashes filenames on every content change.

    A hardcoded asset list would go stale on the next build and publish a page
    pointing at a file that is no longer there.
    """
    from scripts.publish import _referenced_assets

    bundle_assets = [u for u in _referenced_assets() if u.startswith("/static/app/")]

    assert bundle_assets, "the built bundle is not referenced by any page"
    assert any(u.endswith(".js") for u in bundle_assets)
    assert any(u.endswith(".css") for u in bundle_assets)


def test_the_app_route_gets_its_own_shell():
    """Pages routes `/insight` to `insight.html`, so it needs its own copy.

    `explain.html` used to be published from the same bundle. The Calc page was
    deleted from the app, and publishing it here would put a page on Pages that
    the product no longer has."""
    from scripts.publish import APP_ROUTES, APP_SHELL, _published_pages

    pages = _published_pages()

    assert pages["insight.html"] == APP_SHELL
    assert "explain.html" not in pages
    assert "explain" not in APP_ROUTES


def test_the_publisher_refuses_an_incomplete_snapshot():
    """A build-time check, because a silent 200 is not caught by looking."""
    source = (SCRIPTS / "publish.py").read_text(encoding="utf-8")

    assert "def _verify(" in source
    assert "_verify()" in source
    assert "index.html" in source and "200" in source


def test_the_snapshot_refuses_a_tab_pointing_at_a_deleted_route():
    """This used to *rewrite* the `/console` tab into the architecture page,
    because the console POSTed and Pages has no server to POST to. The console
    has since been deleted outright, so there is nothing to rewrite and the
    guard is the whole job: a snapshot must not publish a tab that 404s."""
    import pytest

    from scripts.publish import _snapshot_nav

    for dead in ("/console", "/explain"):
        source = f'<nav class="tabs">\n<a href="{dead}">Gone</a>\n</nav>'
        with pytest.raises(SystemExit) as caught:
            _snapshot_nav(source, "2026-09-07")
        assert dead in str(caught.value)


def test_the_snapshot_says_it_is_a_snapshot_and_when():
    """Presenting frozen data as a live system is the dishonesty this avoids."""
    from scripts.publish import _snapshot_nav

    rewritten = _snapshot_nav('<nav class="rail"><a href="/insight">Insight</a></nav>\n<div class="appbar">\n  <h1>Schedule</h1>\n  <span class="spacer"></span>\n</div>', "2026-09-07")

    assert "static snapshot" in rewritten
    assert "2026-09-07" in rewritten
    assert "runs locally" in rewritten
    # In the app bar, not the rail: navigation is a 64px column now, and a
    # sentence appended inside it would be squeezed to nothing.
    assert rewritten.index("static snapshot") > rewritten.index('appbar')


def test_the_snapshot_note_is_not_added_twice():
    """Publishing twice must not stack banners."""
    from scripts.publish import _snapshot_nav

    once = _snapshot_nav('<nav class="rail"><a href="/insight">Insight</a></nav>\n<div class="appbar">\n  <h1>Schedule</h1>\n  <span class="spacer"></span>\n</div>', "2026-09-07")
    twice = _snapshot_nav(once, "2026-09-08")

    assert twice.count("static snapshot") == 1


def test_published_asset_urls_carry_a_content_hash():
    """A deploy that fixes a broken asset must be visible immediately.

    Pages caches assets for hours. When `/static/shell.css` briefly returned
    `index.html`, every visitor held that HTML and the fixed deploy changed
    nothing they could see - "hard-reload" is an apology, not a fix. Hashing the
    URL makes a republish self-invalidating.
    """
    from scripts.publish import _fingerprint, _version_assets

    static = SCRIPTS.parent / "app" / "api" / "static"
    versioned = _version_assets('<link rel="stylesheet" href="/static/shell.css">')

    digest = _fingerprint(static / "shell.css")
    assert f"/static/shell.css?v={digest}" in versioned
    assert len(digest) == 8


def test_the_hash_changes_only_when_the_bytes_do(tmp_path):
    """Content-addressed, not date-stamped.

    A date stamp would not change on a same-day republish - exactly when a fix
    needs to invalidate.
    """
    from scripts.publish import _fingerprint

    target = tmp_path / "a.css"
    target.write_text("body { color: red }", encoding="utf-8")
    before = _fingerprint(target)

    target.write_text("body { color: red }", encoding="utf-8")
    assert _fingerprint(target) == before

    target.write_text("body { color: blue }", encoding="utf-8")
    assert _fingerprint(target) != before


def test_immutable_caching_is_only_claimed_for_hashed_paths():
    """`immutable` is safe with a content hash and dangerous without one."""
    from scripts.publish import HEADERS

    assert "/static/*" in HEADERS
    assert "immutable" in HEADERS

    source = (SCRIPTS / "publish.py").read_text(encoding="utf-8")
    assert "_version_assets(_snapshot_nav(" in source


# --------------------------------------------------------------------------
# `scripts.replay` deleting a database it was not asked to build
# --------------------------------------------------------------------------


def test_replay_only_deletes_the_database_it_builds():
    """A regression test for a bug that destroyed data.

    `replay` removed `--db` (default `pulse.db`) whatever `DATABASE_URL` said,
    so building any other SQLite database silently deleted the default one. It
    ate a populated `pulse.db` the first time `scripts.serve` seeded a different
    file, and it would have taken a judge's demo data just as readily.
    """
    from pathlib import Path

    from scripts.replay import REPO, sqlite_file_to_remove

    # The case that caused the loss: a different SQLite target.
    assert sqlite_file_to_remove("sqlite:///fresh.db", postgres=False) == (
        REPO / "fresh.db"
    )

    # A server is never a file. This one mapping to a path would be catastrophic
    # rather than merely annoying.
    assert (
        sqlite_file_to_remove("postgresql+psycopg://u:p@host/db", postgres=False)
        is None
    )
    assert sqlite_file_to_remove("sqlite:///pulse.db", postgres=True) is None

    # `sqlite://` with no path is in-memory - there is nothing to unlink.
    assert sqlite_file_to_remove("sqlite://", postgres=False) is None

    # An absolute path is respected rather than re-rooted under the repo.
    absolute = Path("C:/tmp/x.db") if REPO.drive else Path("/tmp/x.db")
    assert sqlite_file_to_remove(
        f"sqlite:///{absolute}", postgres=False
    ) == absolute


def test_the_dead_tab_is_matched_by_href_not_by_label_text():
    """The trap this closes, which outlived the rewrite it was written for.

    The original was a literal `.replace(">Retriever console<", ...)`. Renaming
    that tab turned it into a silent no-op - the snapshot shipped a link that
    could not work - and its test went on passing, because "the old label is
    absent" is trivially true once the old label is gone. Matching the `href`
    is what makes the label free to change, and that still holds now the
    rewrite has become a refusal.
    """
    import pytest

    from scripts.publish import _snapshot_nav

    renamed = '<nav class="tabs">\n<a href="/console">Something Else Entirely</a>\n</nav>'

    with pytest.raises(SystemExit):
        _snapshot_nav(renamed, "2026-09-07")


def test_a_tab_bar_of_live_routes_publishes_untouched():
    """The inverse of the guard, and the case that must not regress into a
    false alarm: every tab points at something Pages serves, so the build has
    nothing to complain about."""
    from scripts.publish import _snapshot_nav

    live = '<nav class="tabs">\n<a href="/insight">Insight</a>\n</nav>'

    assert 'href="/insight"' in _snapshot_nav(live, "2026-09-07")


def test_the_react_pages_carry_no_static_nav_and_pass_through():
    """They build their tab bar in JavaScript, so there is nothing to rewrite.

    Asserted because the guard above must not start failing the build for the
    two pages that legitimately have no `<nav class="tabs">`.
    """
    from scripts.publish import _snapshot_nav

    react_page = '<div id="root"></div><script src="/static/app/x.js"></script>'

    assert "Architecture" not in _snapshot_nav(react_page, "2026-09-07")


def test_an_empty_tab_bar_is_not_treated_as_a_broken_one():
    """A nav with no tabs has no dead console link to worry about.

    This is the distinction the first version of the guard got wrong: it fired
    on any static nav, which broke the two tests that check the freeze stamp.
    """
    from scripts.publish import _snapshot_nav

    assert _snapshot_nav('<nav class="tabs">\n</nav>', "2026-09-07")


def test_a_page_with_no_server_rendered_bar_is_left_alone():
    """The React pages build their chrome in the bundle.

    There is nothing to stamp, so the publisher must say nothing rather than
    inject stray markup into a page it does not understand.
    """
    from scripts.publish import _snapshot_nav

    react_page = '<div id="root"></div>'

    assert _snapshot_nav(react_page, "2026-09-07") == react_page


# --------------------------------------------------------------------------
# `scripts.from_jira_export`: the three ways a Jira export is not a sheet.
#
# Every case here is one this converter met on a real export
# ("Jira Cowork Local 1.xlsx", 17 issues, 421 columns) rather than one imagined
# for a test - which is the same standard the rest of this file is held to.
# --------------------------------------------------------------------------


def _export(tmp_path, header_extra=(), rows=(), preamble=2):
    """A workbook shaped like a Jira "general_report" export."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "general_report"
    for index in range(preamble):
        sheet.append([f"preamble line {index}"])
    sheet.append(["Project", "Key", "Summary", "Status", "Assignee", "Due Date",
                  "Created", *header_extra])
    for row in rows:
        sheet.append(list(row))
    path = tmp_path / "export.xlsx"
    workbook.save(path)
    return path


def test_the_header_row_is_found_below_the_export_preamble(tmp_path):
    """A Jira export opens with the filter name and "Displaying N issues at
    ...", so the table never starts at row 1."""
    from scripts.from_jira_export import read_export

    path = _export(tmp_path, rows=[("P", "P-1", "One", "To Do", "Ann", None, None)])
    headers, rows = read_export(path)
    assert headers[:3] == ["Project", "Key", "Summary"]
    assert len(rows) == 1


def test_only_rows_carrying_a_key_are_issues(tmp_path):
    """One issue is not one row. An export with rich-text Descriptions writes
    each issue across a block, the fields on the first row and the wrapped text
    below - so counting rows, or assuming a fixed stride, invents issues."""
    from scripts.from_jira_export import read_export

    path = _export(
        tmp_path,
        rows=[
            ("P", "P-1", "One", "To Do", "Ann", None, None),
            (None, None, "...wrapped description...", None, None, None, None),
            (None, None, None, None, None, None, None),
            ("P", "P-2", "Two", "Done", "Bo", None, None),
        ],
    )
    _, rows = read_export(path)
    assert len(rows) == 2


def test_a_cell_holding_the_pages_own_javascript_is_dropped(tmp_path):
    """Jira renders some custom fields by shipping the script that draws them,
    and its own rendering failures arrive as prose. A task whose Owner is a
    jQuery call is worse than a task with no Owner."""
    from scripts.from_jira_export import convert, read_export

    path = _export(
        tmp_path,
        rows=[
            ("P", "P-1", "One", "To Do",
             'setTimeout(function(){ $(".x").remove(); })', None, None),
            ("P", "P-2", "Two", "To Do",
             "Error rendering 'aligned-strategy-customfield'.", None, None),
            ("P", "P-3", "Three", "To Do", "Ann", None, None),
        ],
    )
    headers, rows = read_export(path)
    records, _ = convert(headers, rows)
    assert [r["Owner"] for r in records] == [None, None, "Ann"]


def test_a_date_column_keeps_the_day_and_drops_the_time(tmp_path):
    """Jira writes datetimes. This app compares dates, and carrying `19:22`
    into a planned-finish column would suggest the plan is precise to the
    minute when the PM chose a day."""
    from datetime import date, datetime

    from scripts.from_jira_export import convert, read_export

    path = _export(
        tmp_path,
        rows=[("P", "P-1", "One", "To Do", "Ann", datetime(2026, 9, 30, 19, 22), None)],
    )
    headers, rows = read_export(path)
    records, _ = convert(headers, rows)
    assert records[0]["Planned Finish"] == date(2026, 9, 30)


def test_the_output_is_the_blank_template_and_nothing_else(tmp_path):
    """The point of the converter: what it writes has to be the document the
    upload route already accepts. If these drift, a conversion "succeeds" and
    then ingests zero rows, which is the silent failure this whole path has
    been bitten by before."""
    from openpyxl import load_workbook

    from app.ingest.sources.excel.reader import SCHEDULE_CONTRACT, find_sheet
    from scripts.from_jira_export import convert, read_export, write_schedule

    path = _export(
        tmp_path,
        rows=[("P", "P-1", "One", "To Do", "Ann", None, None)],
    )
    headers, rows = read_export(path)
    records, _ = convert(headers, rows)
    out = tmp_path / "schedule.xlsx"
    write_schedule(records, out)

    written = load_workbook(out, read_only=True)
    try:
        first = [c for c in next(written["Activities"].iter_rows(values_only=True))]
    finally:
        written.close()
    assert tuple(first) == SCHEDULE_CONTRACT.template_headers
    # And the reader agrees it is a schedule, which is the claim that matters.
    assert find_sheet(out, SCHEDULE_CONTRACT, preferred="Activities") == "Activities"


def test_a_workbook_with_no_jira_table_is_refused_not_half_converted(tmp_path):
    """`find_sheet` is deliberately forgiving about tab names, so the failure
    to guard against is a workbook of notes converting into an empty schedule
    that ingests silently.

    Raises `NotAJiraExport`, not `SystemExit`: the conversion moved into `app/`
    so the upload route could use it, and a library that kills the process
    would take the API down instead of answering 400. Turning it into an exit
    is the CLI's job, asserted below."""
    from openpyxl import Workbook

    from app.ingest.sources.jira.export_sheet import NotAJiraExport, read_export

    workbook = Workbook()
    workbook.active.append(["some", "notes"])
    path = tmp_path / "notes.xlsx"
    workbook.save(path)

    with pytest.raises(NotAJiraExport):
        read_export(path)


def test_the_cli_turns_a_refusal_into_a_message_not_a_traceback(tmp_path, capsys):
    """The reader is being told their file is the wrong shape, which is not a
    crash - and a traceback would bury the one sentence that helps."""
    import sys

    from openpyxl import Workbook

    from scripts.from_jira_export import main

    workbook = Workbook()
    workbook.active.append(["some", "notes"])
    path = tmp_path / "notes.xlsx"
    workbook.save(path)

    argv = sys.argv
    sys.argv = ["from_jira_export", str(path)]
    try:
        with pytest.raises(SystemExit) as caught:
            main()
    finally:
        sys.argv = argv

    assert "Key" in str(caught.value) and "Summary" in str(caught.value)


def test_the_converter_has_one_implementation_not_two():
    """The CLI and `POST /api/sources/upload` both convert Jira exports. If the
    script grew its own copy they would drift, and the file a person uploads
    would stop matching the file they converted by hand."""
    source = (SCRIPTS / "from_jira_export.py").read_text(encoding="utf-8")

    assert "from app.ingest.sources.jira.export_sheet import" in source
    # The parsing itself must not live here any more.
    assert "def read_export" not in source
    assert "def convert(" not in source

# --------------------------------------------------------------------------
# The columns that are empty today and will not be tomorrow.
#
# Every case below is a way the conversion would have gone wrong *silently* the
# first time somebody filled a column in - no error, just a worse answer.
# --------------------------------------------------------------------------


def _future_export(tmp_path, header, rows):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "general_report"
    sheet.append(["FI2.0"])
    sheet.append(["Displaying N issues at 13/Sep/26 9:00 AM."])
    sheet.append(header)
    for row in rows:
        sheet.append(list(row))
    path = tmp_path / "future.xlsx"
    workbook.save(path)
    return path


LINK_HEADER = ["Project", "Key", "Summary", "Status", "Linked Issues", "Linked Issues"]


def _converted(tmp_path, header, rows):
    from app.ingest.sources.jira.export_sheet import convert, read_export

    headers, data = read_export(_future_export(tmp_path, header, rows))
    records, chosen = convert(headers, data)
    return {r["Task ID"]: r for r in records}, chosen


def test_a_repeated_header_is_read_across_every_column(tmp_path):
    """Jira repeats a header per value rather than packing a list into one cell
    - this export already carries `Approver` four times. Keeping one position
    per header reads only the last, so an issue with two links would contribute
    one edge and lose the other."""
    by, _ = _converted(
        tmp_path, LINK_HEADER,
        [("P", "P-2", "Build", "To Do", "depends on P-1", "is blocked by P-3")],
    )
    assert by["P-2"]["Predecessor"] == "P-1, P-3"


def test_an_outbound_link_never_becomes_a_predecessor(tmp_path):
    """"blocks P-2" and "is blocked by P-2" name the same pair and opposite
    edges. Accepting both reverses half the arrows, which does not fail - it
    returns a confident, wrong critical path."""
    by, _ = _converted(
        tmp_path, LINK_HEADER,
        [("P", "P-3", "Spec", "To Do", "blocks P-2", None)],
    )
    assert by["P-3"]["Predecessor"] is None


def test_a_bare_key_with_no_direction_is_dropped_not_guessed(tmp_path):
    """Same rule `dependencies.py` applies to a dangling reference: drop it
    rather than hedge. A key on its own does not say which way the arrow goes."""
    by, _ = _converted(
        tmp_path, LINK_HEADER, [("P", "P-4", "Infra", "To Do", "P-1", None)]
    )
    assert by["P-4"]["Predecessor"] is None


def test_a_sub_task_is_never_read_as_a_predecessor(tmp_path):
    """A sub-task is a child. Feeding it to `Predecessor` would assert that a
    parent waits for its children - a chain nobody stated, reaching a critical
    path and a projected date."""
    by, _ = _converted(
        tmp_path,
        ["Project", "Key", "Summary", "Status", "Sub-Tasks"],
        [("P", "P-1", "Design", "To Do", "P-9")],
    )
    assert by["P-1"]["Predecessor"] is None


@pytest.mark.parametrize(
    "raw,expected",
    [("100%", 100.0), (60, 60.0), ("60 %", 60.0), (43200, None), ("blocks P-2", None)],
)
def test_progress_is_taken_as_a_percentage_or_not_at_all(tmp_path, raw, expected):
    """43200 is Jira's aggregate progress in *seconds*. Rescaling it would put a
    plausible completion figure on screen; rejecting it says nothing, which is
    what we know."""
    by, _ = _converted(
        tmp_path,
        ["Project", "Key", "Summary", "Status", "Progress"],
        [("P", "P-1", "Design", "To Do", raw)],
    )
    assert by["P-1"]["Progress"] == expected


def test_an_outbound_only_column_says_so_rather_than_reporting_no_links(tmp_path):
    """"You have links, they point the other way" and "you have no links" need
    different actions, and only one of them is fixable in Jira in a minute."""
    from app.ingest.sources.jira.export_sheet import convert, coverage, read_export

    headers, data = read_export(
        _future_export(tmp_path, LINK_HEADER, [("P", "P-3", "Spec", "To Do", "blocks P-2", None)])
    )
    records, chosen = convert(headers, data)
    notes = " ".join(coverage(records, chosen, headers, data)["notes"])
    assert "outbound direction" in notes

def test_a_populated_but_meaningless_column_loses_to_a_meaningful_one(tmp_path):
    """"Not empty" is not enough to identify a column.

    Jira defines every standard field whether or not it holds anything useful,
    so an export can carry a populated `Parent Link` holding a documentation URL
    beside a populated `Product` holding the real `Name [KEY]`. Ranking on
    non-emptiness alone picked the URL and made it a milestone name.
    """
    by, chosen = _converted(
        tmp_path,
        ["Project", "Key", "Summary", "Status", "Parent Link", "Product"],
        [("P", "P-1", "Design", "To Do", "https://example.com/docs", "Platform [P-100]")],
    )
    assert chosen["Milestone"] == "Product"
    assert by["P-1"]["Milestone"] == "Platform"


def test_a_url_is_never_a_milestone_name(tmp_path):
    """Even when it is the only candidate. A link to a thing is not the name of
    one, and a grouping band labelled `https://...` is worse than an unlabelled
    one."""
    by, _ = _converted(
        tmp_path,
        ["Project", "Key", "Summary", "Status", "Parent Link"],
        [("P", "P-1", "Design", "To Do", "https://example.com/docs")],
    )
    assert by["P-1"]["Milestone"] is None


def test_a_column_that_merely_copies_another_loses_to_a_real_one(tmp_path):
    """`Planned Start` is stamped at creation on an instance that does not use
    it, so it equals `Created` on every row.

    The coverage report already detected that and said so - while the conversion
    went on preferring it over a populated `Start date` holding real, different
    dates. Detecting a field default and then choosing it anyway is worse than
    not detecting it.
    """
    from datetime import datetime

    created = datetime(2026, 9, 2, 19, 22)
    by, chosen = _converted(
        tmp_path,
        ["Project", "Key", "Summary", "Status", "Created", "Planned Start", "Start date"],
        [("P", "P-1", "Design", "To Do", created, created, datetime(2026, 9, 17))],
    )
    assert chosen["Start"] == "Start date"
    assert by["P-1"]["Start"].isoformat() == "2026-09-17"


def test_the_only_candidate_is_still_used_even_if_it_copies_another(tmp_path):
    """Demotion is a preference between candidates, not a veto.

    With no `Start date` to fall back on, `Planned Start` is all there is - and
    the coverage note saying it is a field default is the honest way to serve it,
    which is what the real export already does.
    """
    from datetime import datetime

    created = datetime(2026, 9, 2, 19, 22)
    by, chosen = _converted(
        tmp_path,
        ["Project", "Key", "Summary", "Status", "Created", "Planned Start"],
        [("P", "P-1", "Design", "To Do", created, created)],
    )
    assert chosen["Start"] == "Planned Start"
    assert by["P-1"]["Start"].isoformat() == "2026-09-02"

# --------------------------------------------------------------------------
# Hostile input. Every case here crashed or lied before it was fixed.
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "status,expect_open,expect_done",
    [
        ("Done", False, True),
        ("Closed", False, True),
        # The one that mattered: `Resolved` is among the commonest states in a
        # real Jira workflow and fell through to OTHER, which every count reads
        # as open - so a finished task was reported overdue, in progress and
        # stale at once.
        ("Resolved", False, True),
        # Work that will not happen is neither delivered nor outstanding.
        # Counting it done inflates completion; counting it open reports a
        # cancelled task as late forever.
        ("Cancelled", False, False),
        ("Won't Do", False, False),
        ("Rejected", False, False),
        ("In Progress", True, False),
        ("Backlog", True, False),
        # Genuinely unknown stays unknown rather than being guessed at.
        ("Sausages", True, False),
    ],
)
def test_a_real_jira_status_is_classified_not_guessed(status, expect_open, expect_done):
    from app.ingest.sources.excel.convertor import _normalize_status
    from app.intelligence.context import CLOSED_STATES, DONE_STATES

    normalised = _normalize_status(status).casefold()
    assert (normalised not in CLOSED_STATES) is expect_open, status
    assert (normalised in DONE_STATES) is expect_done, status


def test_a_date_at_the_edge_of_the_calendar_does_not_take_a_page_down():
    """A tracker can hold 9999-12-31 as a "no due date" sentinel, and it is a
    trivial thing to mistype. `date.max + timedelta(days=1)` raises rather than
    saturating, so the schedule window and the due-soon window both 500'd on
    one such row."""
    from datetime import date

    from app.intelligence.context import shift_date

    assert shift_date(date.max, 3) == date.max
    assert shift_date(date.min, -3) == date.min
    assert shift_date(date(2026, 9, 12), 3) == date(2026, 9, 15)


def test_a_hostile_export_converts_without_raising(tmp_path):
    """One row per way to fool the converter. None of them may raise, and none
    may become a confident wrong value."""
    from datetime import datetime

    from app.ingest.sources.jira.export_sheet import convert, read_export

    header = ["Project", "Key", "Summary", "Status", "Due Date", "Start date",
              "Progress", "Linked Issues"]
    rows = [
        ("T", "T-1", "text in a date", "To Do", "TBD", "N/A", None, None),
        ("T", "T-2", "far future", "To Do", datetime(9999, 12, 31), None, None, None),
        ("T", "T-3", "progress over 100", "To Do", None, None, 9999, None),
        ("T", "T-4", "negative progress", "To Do", None, None, -50, None),
        ("T", "T-5", "self dependency", "To Do", None, None, None, "is blocked by T-5"),
        ("T", "T-6", "x" * 5000, "To Do", None, None, None, None),
        ("T", "T-7", "unicode 工数 😀", "To Do", None, None, None, None),
    ]
    by, _ = _converted(tmp_path, header, rows)

    assert by["T-1"]["Planned Finish"] is None and by["T-1"]["Start"] is None
    assert by["T-2"]["Planned Finish"].year == 9999
    assert by["T-3"]["Progress"] is None
    assert by["T-4"]["Progress"] is None
    # Kept as written - dropping a self-edge is `dependencies.py`'s job, and it
    # does it. The converter must not silently rewrite what the source said.
    assert by["T-5"]["Predecessor"] == "T-5"
    assert len(by["T-6"]["Activity"]) == 5000
    assert "😀" in by["T-7"]["Activity"]
