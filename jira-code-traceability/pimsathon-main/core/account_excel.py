"""Bulk-create Groups + Accounts from an Excel list (Admin only).

Mirrors ``task_excel.py``'s template/import pattern: the app exports a
ready-made ``.xlsx`` template (with a README sheet), the Admin fills one row
per person, and importing creates any missing Groups by name plus one
Account per row — each with a freshly generated 12-character access code.
The generated codes are returned so the caller can show/export them (they
exist nowhere else in plain sight; the Admin hands them out).
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional, Tuple

from . import accounts, groups

HEADERS = ["Username", "Display name", "Email", "Department", "Role", "Group"]
_EXAMPLE_ROWS = [
    ["nguyenva1", "Nguyen Van A", "nguyenva1@company.com", "CAE", "user", "CAE Team"],
    ["tranthib2", "Tran Thi B", "tranthib2@company.com", "CAE", "subadmin", "CAE Team"],
]


def export_template(path: str | Path) -> Path:
    """Write the import template (Accounts sheet + README) to ``path``."""
    from openpyxl import Workbook

    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"
    ws.append(HEADERS)
    for row in _EXAMPLE_ROWS:
        ws.append(row)
    for col, width in zip("ABCDEF", (18, 24, 26, 16, 12, 20)):
        ws.column_dimensions[col].width = width

    readme = wb.create_sheet("README")
    for line in (
        "One row per person. Username: lowercase letters/digits (the login name).",
        "Role: user or subadmin. 'admin' is NOT allowed here — the app has exactly one Admin.",
        "Group: a group name; groups that don't exist yet are created automatically.",
        "On import, every account gets a fresh 12-character access code —",
        "the app shows/exports the codes once so the Admin can distribute them.",
        "Rows whose Username already exists are skipped (never overwritten).",
    ):
        readme.append([line])
    readme.column_dimensions["A"].width = 100

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def import_accounts(path: str | Path, shared_dir: str,
                    created_by: str = "") -> Tuple[List[accounts.Account], List[str]]:
    """Create groups + accounts from a filled template.

    Returns ``(created_accounts, warnings)``. Never creates a second admin
    (role 'admin' rows are downgraded to 'user' with a warning), never
    overwrites an existing username (skipped with a warning). Raises
    ``ValueError`` for an unreadable/empty file — mirroring
    ``task_excel.import_tasks``'s error contract."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(str(path), data_only=True)
    except Exception as exc:  # noqa: BLE001 — one readable error for any bad file
        raise ValueError(f"Could not read the Excel file: {exc}") from exc
    ws = wb["Accounts"] if "Accounts" in wb.sheetnames else wb.worksheets[0]

    acc_dir = accounts.accounts_dir(shared_dir)
    grp_dir = groups.groups_dir(shared_dir)
    existing_codes = {a.code for a in accounts.list_accounts(acc_dir)}
    groups_by_name: Dict[str, groups.Group] = {
        g.name.strip().lower(): g for g in groups.list_groups(grp_dir)
    }

    created: List[accounts.Account] = []
    warnings: List[str] = []
    rows = ws.iter_rows(min_row=2, values_only=True)
    for i, row in enumerate(rows, start=2):
        cells = [str(c).strip() if c is not None else "" for c in (row or ())]
        cells += [""] * (len(HEADERS) - len(cells))
        username, display_name, email, department, role, group_name = cells[:6]
        if not username:
            continue
        role = (role or "user").lower()
        if role == "admin":
            warnings.append(f"Row {i}: role 'admin' is not allowed (single-admin app) — created as 'user'.")
            role = "user"
        if role not in accounts.ROLES:
            warnings.append(f"Row {i}: unknown role '{role}' — created as 'user'.")
            role = "user"
        if accounts.load_account(username, acc_dir) is not None:
            warnings.append(f"Row {i}: account '{username}' already exists — skipped.")
            continue

        group_id = ""
        if group_name:
            key = group_name.strip().lower()
            group = groups_by_name.get(key)
            if group is None:
                group = groups.new_group(group_name.strip())
                groups.save_group(group, grp_dir)
                groups_by_name[key] = group
            group_id = group.group_id

        account = accounts.new_account(
            username, role, display_name=display_name, department=department,
            email=email, group_id=group_id, created_by=created_by,
            existing_codes=existing_codes)
        accounts.save_account(account, acc_dir)
        existing_codes.add(account.code)
        created.append(account)

        if group_id:
            group = groups_by_name[group_name.strip().lower()]
            if role == "subadmin" and not group.subadmin_username:
                group.subadmin_username = account.username
                groups.save_group(group, grp_dir)
            elif account.username not in group.member_usernames:
                group.member_usernames.append(account.username)
                groups.save_group(group, grp_dir)

    if not created and not warnings:
        raise ValueError("No account rows found in the file (fill the 'Accounts' sheet).")
    return created, warnings


def export_issued_codes(created: List[accounts.Account], path: str | Path) -> Optional[Path]:
    """Write the just-created accounts + their access codes to an xlsx the
    Admin can distribute from. Best-effort: returns None on write failure
    (the codes were already shown in the UI)."""
    from openpyxl import Workbook

    if not created:
        return None
    try:
        path = Path(path)
        wb = Workbook()
        ws = wb.active
        ws.title = "Issued codes"
        ws.append(["Username", "Display name", "Role", "Group", "Access code"])
        for a in created:
            ws.append([a.username, a.display_name, a.role, a.group_id, a.code])
        for col, width in zip("ABCDE", (18, 24, 12, 20, 18)):
            ws.column_dimensions[col].width = width
        wb.save(path)
        return path
    except Exception:  # noqa: BLE001
        return None
