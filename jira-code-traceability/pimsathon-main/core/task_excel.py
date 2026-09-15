"""Schedule Task — Excel template export + task import.

``export_template(path)`` writes an .xlsx the user fills in;
``import_tasks(path)`` turns its rows back into task dicts (NOT yet saved —
the Import tab previews them and only saves after the user confirms).

The "Depends on" column references OTHER ROWS' Title values (semicolon-
separated) so a whole parallel fan-in graph can be described in one file:
those references resolve to real task ids after all rows are created.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from .tasks import (
    PRIORITIES, REPEAT_TYPES, TASK_TYPES, new_task, parse_run_at,
)

HEADERS = [
    "Title", "Description", "Type", "Priority", "Script command",
    "Schedule enabled", "Run at (YYYY-MM-DD HH:MM)", "Repeat",
    "Cron expression", "Depends on (titles, ;-separated)",
    "Use previous output as input", "Requires approval",
    # Newer optional columns — blank falls back to the app's Settings default.
    "Provider", "Model", "Skill (slug)",
    "Reminder (none/teams/outlook)", "Reminder email",
]

_EXAMPLE = [
    "Generate CAE report", "Đọc dữ liệu CAE mới và tạo báo cáo markdown",
    "co4e_code", "high", "", "yes", "2026-07-06 09:00", "weekly", "",
    "", "no", "no", "", "", "", "none", "",
]
_EXAMPLE2 = [
    "Draft team email", "Soạn email draft từ báo cáo",
    "cowork", "medium", "", "no", "", "none", "",
    "Generate CAE report", "yes", "no", "", "", "", "outlook", "boss@fpt.com",
]

_TRUE = {"yes", "y", "true", "1", "x", "có", "co"}


def _bool(value) -> bool:
    return str(value or "").strip().lower() in _TRUE


def _clamp(value, allowed, default):
    v = str(value or "").strip().lower()
    return v if v in allowed else default


def export_template(path: str | Path) -> Path:
    """Write the fill-in template (headers + 2 linked example rows + notes).

    Every column whose value is one of a fixed set (Type, Priority, Repeat, the
    yes/no flags, Reminder channel) gets an in-cell DROPDOWN so the user just
    picks a valid value instead of typing it — fewer import errors."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="F37021")
    ws.append(_EXAMPLE)
    ws.append(_EXAMPLE2)
    for col, header in enumerate(HEADERS, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = \
            max(18, len(header) + 2)

    # In-cell dropdowns for the list/enum columns (1-indexed to HEADERS order).
    _YESNO = ["yes", "no"]
    dropdowns = {
        3: list(TASK_TYPES),           # Type
        4: list(PRIORITIES),           # Priority
        6: _YESNO,                     # Schedule enabled
        8: list(REPEAT_TYPES),         # Repeat
        11: _YESNO,                    # Use previous output as input
        12: _YESNO,                    # Requires approval
        16: ["none", "teams", "outlook"],   # Reminder channel
    }
    for col_idx, options in dropdowns.items():
        letter = ws.cell(row=1, column=col_idx).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                            allow_blank=True, showDropDown=False)
        dv.prompt = "Chọn một giá trị từ danh sách"
        dv.error = "Giá trị không hợp lệ — chọn từ danh sách."
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}500")   # apply to the fill-in rows

    notes = wb.create_sheet("README")
    notes.append(["Điền mỗi task một dòng trong sheet 'Tasks' (2 dòng ví dụ có sẵn — xoá hoặc sửa)."])
    notes.append([f"Type: {', '.join(TASK_TYPES)}   ·   Priority: {', '.join(PRIORITIES)}"])
    notes.append([f"Repeat: {', '.join(REPEAT_TYPES)} (cron → điền thêm cột Cron expression)"])
    notes.append(["Depends on: tên (Title) các dòng khác, cách nhau dấu ';' — task này chỉ chạy khi các task đó Done."])
    notes.append(["Use previous output as input = yes → output các task Depends-on tự thành input task này."])
    path = Path(path)
    wb.save(str(path))
    return path


def task_from_cells(cells: List[Any]):
    """Turn one row (a list in HEADERS order) into ``(task, depends_titles)`` or
    ``(None, None)`` for a blank row. Shared by the Excel and CSV importers so
    both apply exactly the same field mapping and enum fallbacks."""
    cells = list(cells) + [None] * (len(HEADERS) - len(cells))
    title = str(cells[0] or "").strip()
    if not title:
        return None, None
    t = new_task(title)
    t["description"] = str(cells[1] or "").strip()
    t["task_type"] = _clamp(cells[2], TASK_TYPES, "manual")
    t["priority"] = _clamp(cells[3], PRIORITIES, "medium")
    t["script_command"] = str(cells[4] or "").strip()
    t["schedule"]["enabled"] = _bool(cells[5])
    run_at = str(cells[6] or "").strip()[:16]
    t["schedule"]["run_at"] = run_at if parse_run_at(run_at) else None
    t["schedule"]["repeat_type"] = _clamp(cells[7], REPEAT_TYPES, "none")
    t["schedule"]["cron_expression"] = str(cells[8] or "").strip() or None
    if t["schedule"]["enabled"] and (t["schedule"]["run_at"]
                                     or t["schedule"]["repeat_type"] == "cron"):
        t["status"] = "scheduled"
    else:
        t["schedule"]["enabled"] = False   # no usable time → stays Backlog
    if _bool(cells[10]):
        t["input"]["mode"] = "previous_task_output"
    t["execution"]["requires_approval"] = _bool(cells[11])
    # Newer optional columns (blank → keep new_task defaults / Settings model).
    t["provider"] = str(cells[12] or "").strip()
    t["model"] = str(cells[13] or "").strip()
    t["skill_slug"] = str(cells[14] or "").strip()
    channel = str(cells[15] or "").strip().lower()
    t["execution"]["notify_channel"] = channel if channel in ("teams", "outlook") else "none"
    t["execution"]["notify_email"] = str(cells[16] or "").strip()
    if t["execution"]["notify_channel"] != "none":
        t["execution"]["notify_on_complete"] = True
        t["execution"]["notify_on_error"] = True
    depends = [s.strip() for s in str(cells[9] or "").split(";") if s.strip()]
    return t, depends


def resolve_depends(tasks: List[Dict[str, Any]], depends_raw: List[List[str]]) -> None:
    """Resolve each row's "Depends on" titles → ids (rows in this same file
    only) and wire output-passing for previous-output inputs. Mutates in place."""
    by_title = {t["title"]: t for t in tasks}
    for t, wants in zip(tasks, depends_raw):
        ids = [by_title[w]["task_id"] for w in wants if w in by_title
               and by_title[w] is not t]
        t["dependency"]["depends_on"] = ids
        if ids and t["input"]["mode"] == "previous_task_output":
            for pid in ids:   # predecessors pass their output forward
                by_id_task = next(x for x in tasks if x["task_id"] == pid)
                by_id_task["dependency"]["pass_output_to_next"] = True


def import_tasks(path: str | Path) -> List[Dict[str, Any]]:
    """Parse a filled template into task dicts with dependencies resolved.
    Raises ValueError with a human message on unusable files; skips blank
    rows; bad enum cells fall back to safe defaults instead of failing."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(str(path), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Cannot read Excel file: {exc}") from exc
    ws = wb["Tasks"] if "Tasks" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    tasks: List[Dict[str, Any]] = []
    depends_raw: List[List[str]] = []
    for row in rows:
        t, depends = task_from_cells(list(row))
        if t is None:
            continue
        depends_raw.append(depends)
        tasks.append(t)

    if not tasks:
        raise ValueError("No task rows found — fill in the 'Tasks' sheet first.")
    resolve_depends(tasks, depends_raw)
    return tasks
