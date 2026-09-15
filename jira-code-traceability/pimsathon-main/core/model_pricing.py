"""LLM model price list — a rich, importable/exportable table shown in the
Monitoring Overview.

Each entry mirrors the vendor price sheet columns:
    Model name · Context length · Max output token · Input price · Input Unit ·
    Output price · Output Unit

Prices carry their own currency (parsed from the ₫ / ¥ / $ symbol, or the
config default) and are converted to the display currency (VND / JPY / USD)
using the same USD↔VND↔JPY rates the usage tracker uses.

Stored in the app config under ``model_pricing.entries`` so it persists and can
be edited by hand, imported from a template, or auto-linked from the providers.
"""
from __future__ import annotations

import copy
import csv
from pathlib import Path
from typing import Any, Dict, List, Optional

# Template columns, in order (exact vendor-sheet layout).
COLUMNS = ["Model name", "Context length", "Max output token",
           "Input price", "Input Unit", "Output price", "Output Unit"]

_EXAMPLE_ROW = ["FPT.AI-KIE-v1.7", "33k", "33k",
                "20.359 ₫", "Million tokens", "20.359 ₫", "Million tokens"]

_SYMBOL_CCY = {"₫": "VND", "vnd": "VND", "đ": "VND",
               "¥": "JPY", "jpy": "JPY", "yen": "JPY",
               "$": "USD", "usd": "USD"}

_DEFAULT_UNIT = "Million tokens"


# ---- currency ------------------------------------------------------------
def _rates(config) -> Dict[str, float]:
    """1 USD = X <currency>. Reuses the usage tracker's editable rates."""
    usage = (getattr(config, "data", {}) or {}).get("usage", {}) if config else {}
    return {"USD": 1.0,
            "VND": float(usage.get("usd_to_vnd", 25000.0) or 25000.0),
            "JPY": float(usage.get("usd_to_jpy", 150.0) or 150.0)}


def convert(amount: float, from_ccy: str, to_ccy: str, config=None) -> float:
    """Convert ``amount`` from one supported currency to another."""
    rates = _rates(config)
    frm = rates.get((from_ccy or "USD").upper(), 1.0)
    to = rates.get((to_ccy or "USD").upper(), 1.0)
    if frm <= 0:
        return amount
    return amount / frm * to


_SYMBOLS = {"VND": "₫", "JPY": "¥", "USD": "$"}
_DIGITS = {"VND": 0, "JPY": 1, "USD": 4}


def format_price(amount: float, ccy: str) -> str:
    ccy = (ccy or "USD").upper()
    return f"{amount:,.{_DIGITS.get(ccy, 2)}f} {_SYMBOLS.get(ccy, '')}".strip()


def parse_price(text: Any) -> tuple:
    """Parse a price cell like ``"20.359 ₫"`` / ``"$0.15"`` → ``(amount, ccy)``.
    ``ccy`` is ``None`` when no symbol is present (caller supplies a default).
    VND is treated as integer thousands (``20.359`` → ``20359``); other
    currencies use ``.`` as the decimal separator."""
    if text is None:
        return 0.0, None
    s = str(text).strip()
    if not s:
        return 0.0, None
    ccy = None
    low = s.lower()
    for sym, code in _SYMBOL_CCY.items():
        if sym in low:
            ccy = code
            break
    # strip everything but digits and separators
    cleaned = "".join(ch for ch in s if ch.isdigit() or ch in ".,")
    if not cleaned:
        return 0.0, ccy
    try:
        if ccy == "VND":
            return float(cleaned.replace(".", "").replace(",", "")), ccy
        return float(cleaned.replace(",", "")), ccy
    except ValueError:
        return 0.0, ccy


# ---- store ---------------------------------------------------------------
def _bucket(config) -> Dict[str, Any]:
    return config.data.setdefault("model_pricing", {})


def list_entries(config) -> List[Dict[str, Any]]:
    return list(_bucket(config).get("entries", []) or [])


def save_entries(config, entries: List[Dict[str, Any]]) -> None:
    _bucket(config)["entries"] = [dict(e) for e in entries]
    sync_to_usage(config)      # keep the cost engine (Overview + Dashboard) in sync


def _norm_entry(model: str, ctx_len: str = "", max_out: str = "",
                in_price=0.0, in_ccy: Optional[str] = None, in_unit: str = _DEFAULT_UNIT,
                out_price=0.0, out_ccy: Optional[str] = None, out_unit: str = _DEFAULT_UNIT,
                default_ccy: str = "USD") -> Dict[str, Any]:
    return {
        "model": str(model).strip(),
        "context_length": str(ctx_len).strip(),
        "max_output": str(max_out).strip(),
        "input_price": float(in_price or 0.0),
        "input_ccy": (in_ccy or default_ccy).upper(),
        "input_unit": (in_unit or _DEFAULT_UNIT).strip(),
        "output_price": float(out_price or 0.0),
        "output_ccy": (out_ccy or default_ccy).upper(),
        "output_unit": (out_unit or _DEFAULT_UNIT).strip(),
    }


def usd_rates_for(model: str, config) -> Optional[Dict[str, float]]:
    """USD price per 1M tokens for ``model`` from the price table (converting the
    entry's own currency to USD), or None when the model isn't in the table."""
    for e in list_entries(config):
        if e.get("model") == model:
            return {
                "in": convert(float(e.get("input_price", 0) or 0), e.get("input_ccy", "USD"), "USD", config),
                "out": convert(float(e.get("output_price", 0) or 0), e.get("output_ccy", "USD"), "USD", config),
            }
    return None


def turn_cost_usd(model: str, in_tok: int, out_tok: int, config) -> float:
    """Cost (USD) of a turn — uses the model's row in the price table when present,
    else the usage tracker's flat fallback rates. Auto-updates when the user
    switches models (a different model → its own row / rates)."""
    rates = usd_rates_for(model, config)
    if rates is None:
        from . import usage_tracker as ut
        p = {**ut.DEFAULT_PRICING, **((getattr(config, "data", {}) or {}).get("usage") or {})}
        rates = {"in": float(p["price_per_mtok_in_usd"]), "out": float(p["price_per_mtok_out_usd"])}
    return (in_tok or 0) / 1e6 * rates["in"] + (out_tok or 0) / 1e6 * rates["out"]


def sync_to_usage(config) -> Dict[str, Dict[str, float]]:
    """Push this table's per-model USD rates into the usage tracker's
    ``usage.model_prices`` map, so token-cost TOTALS on the Monitoring Overview
    cards AND the Dashboard chart are computed straight from THIS price table
    (and update automatically whenever it is imported/edited). Only rows that
    carry a non-zero price are pushed; unpriced models fall back to the flat
    ``price_per_mtok_*`` rates in the usage tracker."""
    if config is None:
        return {}
    usage = config.data.setdefault("usage", {})
    table: Dict[str, Dict[str, float]] = {}
    for e in list_entries(config):
        model = str(e.get("model", "")).strip()
        if not model:
            continue
        in_usd = convert(float(e.get("input_price", 0) or 0), e.get("input_ccy", "USD"), "USD", config)
        out_usd = convert(float(e.get("output_price", 0) or 0), e.get("output_ccy", "USD"), "USD", config)
        if in_usd <= 0 and out_usd <= 0:
            continue        # unpriced row → leave this model to the flat fallback
        table[model] = {"in": in_usd, "out": out_usd}
    usage["model_prices"] = table
    return table


def format_tokens(n: int) -> str:
    """Compact token count: 108 · 2.0k · 104.8k · 3.29M."""
    n = int(n or 0)
    if n >= 1_000_000:
        return f"{n / 1e6:.2f}M"
    if n >= 1000:
        return f"{n / 1000:.1f}k"
    return str(n)


def add_entry(config, entry: Dict[str, Any]) -> None:
    entries = list_entries(config)
    entries = [e for e in entries if e.get("model") != entry.get("model")]  # replace same model
    entries.append(entry)
    save_entries(config, entries)


def entry_from_row(cells: List[Any], default_ccy: str = "USD"):
    """Build an entry from a template row (list in COLUMNS order). Returns None
    for a blank/header row."""
    cells = list(cells) + [None] * (len(COLUMNS) - len(cells))
    model = str(cells[0] or "").strip()
    if not model or model.lower() == COLUMNS[0].lower():
        return None
    in_amt, in_ccy = parse_price(cells[3])
    out_amt, out_ccy = parse_price(cells[5])
    return _norm_entry(model, cells[1] or "", cells[2] or "",
                       in_amt, in_ccy, str(cells[4] or _DEFAULT_UNIT),
                       out_amt, out_ccy, str(cells[6] or _DEFAULT_UNIT),
                       default_ccy=default_ccy)


# ---- import / export -----------------------------------------------------
def export_template(path: str | Path) -> Path:
    """Write the fill-in price template (headers + 1 example row) as .xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Pricing"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="F37021")
    ws.append(_EXAMPLE_ROW)
    for col, header in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = max(16, len(header) + 2)
    path = Path(path)
    wb.save(str(path))
    return path


def import_table(path: str | Path, default_ccy: str = "USD") -> List[Dict[str, Any]]:
    """Parse a filled template (.xlsx / .csv) into entries. Raises ValueError on
    an unusable file; blank rows are skipped."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xls", ".xlsm"):
        rows = _rows_from_xlsx(p)
    elif ext == ".csv":
        rows = _rows_from_csv(p)
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Use .xlsx or .csv.")
    entries = []
    for r in rows:
        e = entry_from_row(r, default_ccy=default_ccy)
        if e is not None:
            entries.append(e)
    if not entries:
        raise ValueError("No price rows found — fill in the template first.")
    return entries


def _rows_from_xlsx(path: Path) -> List[List[Any]]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(str(path), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Cannot read Excel file: {exc}") from exc
    ws = wb["Pricing"] if "Pricing" in wb.sheetnames else wb.active
    return [list(r) for r in ws.iter_rows(min_row=1, values_only=True)]


def _rows_from_csv(path: Path) -> List[List[Any]]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except OSError as exc:
        raise ValueError(f"Cannot read CSV file: {exc}") from exc
    return [list(r) for r in csv.reader(text.splitlines())]


# ---- auto-link from the providers ---------------------------------------
def auto_link(ctx, default_ccy: str = "USD") -> List[Dict[str, Any]]:
    """Fetch the live model list from the configured providers and add a row for
    each NEW model (blank prices, to be filled in). Returns the merged list and
    saves it. Best-effort: unreachable providers are simply skipped."""
    from . import preview_ai

    try:
        by_provider = preview_ai.fetch_live_models(ctx) or {}
    except Exception:  # noqa: BLE001
        by_provider = {}
    models = []
    for lst in by_provider.values():
        models.extend(lst or [])
    entries = list_entries(ctx.config)
    have = {e.get("model") for e in entries}
    for m in sorted(set(models)):
        if m and m not in have:
            entries.append(_norm_entry(m, default_ccy=default_ccy))
            have.add(m)
    save_entries(ctx.config, entries)
    return entries
